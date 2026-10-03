"""Find trades on Google Maps - sole traders and partnerships included, not just limited companies.

Called by find_prospects.py --source google (the panel's finder, "Google Maps" source). Companies
House only lists limited companies; most roofers, landscapers and builders are sole traders, and
Google Maps has them - with their website, phone and how many reviews they have.

    For every trade x area: Google Places "Text Search" ("roofer in Guildford"), up to 60 results.
    Kept: open businesses that aren't already in one of your lists, chains and merchants left out.
    Busiest first (most Google reviews): established firms that clearly get work, and can afford it.
    Their website goes on the Outreach tab; no website (or only a Facebook / Checkatrade page) goes
    on the No website tab, for a letter. Emails aren't on Google - Run the whole list finds them
    on their sites, as it does for any list.

Needs GOOGLE_PLACES_API_KEY in Settings, or uses PAGESPEED_API_KEY if it's from the same Google
Cloud project - either way, enable "Places API (New)" on that project. Google charges per search
(each returns up to 20 firms) after a monthly free allowance; a 100-firm list is a handful of
searches.
"""

from __future__ import annotations

import json
import re
import urllib.error
import urllib.request
from pathlib import Path

ENDPOINT = "https://places.googleapis.com/v1/places:searchText"
FIELDS = ",".join(["places.id", "places.displayName", "places.websiteUri", "places.nationalPhoneNumber",
                   "places.formattedAddress", "places.rating", "places.userRatingCount", "places.businessStatus",
                   "places.primaryType", "nextPageToken"])
PAGES = 3  # Google stops at 60 results a search
SEARCH_WORDS = {
    "lofts": "loft conversion company",
    "driveways": "driveway contractor",
    "roofing": "roofer",
    "landscaping": "landscaper",
    "building": "builder",
}
# Merchants, hire shops and chains that turn up for "builder" - never a prospect.
NOT_A_PROSPECT = {"hardware_store", "home_improvement_store", "home_goods_store", "store", "furniture_store",
                  "real_estate_agency", "lodging", "shopping_mall", "warehouse_store"}
CHAINS = re.compile(r"\b(wickes|b&q|travis perkins|jewson|screwfix|toolstation|selco|homebase|howdens|buildbase|"
                    r"huws gray|ibstock|marshalls|topps tiles|speedy hire|hss hire|checkatrade|mybuilder|rated people)\b", re.I)
EXTRA_HEADERS = ["Google rating", "Google reviews", "Source"]


class PlacesError(Exception):
    pass


def search(query: str, key: str, post=None) -> list[dict]:
    """Every result for one text search (up to 60), as Google returns them."""
    post = post or _post
    out, token = [], ""
    for _ in range(PAGES):
        body = {"textQuery": query, "regionCode": "gb", "languageCode": "en-GB", "pageSize": 20}
        if token:
            body["pageToken"] = token
        res = post(body, key)
        out += res.get("places") or []
        token = res.get("nextPageToken") or ""
        if not token:
            break
    return out


def _post(body: dict, key: str) -> dict:
    req = urllib.request.Request(ENDPOINT, data=json.dumps(body).encode(), method="POST",
                                 headers={"Content-Type": "application/json", "X-Goog-Api-Key": key, "X-Goog-FieldMask": FIELDS})
    try:
        with urllib.request.urlopen(req, timeout=30) as res:
            return json.loads(res.read())
    except urllib.error.HTTPError as e:
        detail = e.read(500).decode("utf-8", errors="replace")
        if e.code in (400, 403) and ("not been used" in detail or "disabled" in detail or "PERMISSION_DENIED" in detail):
            raise PlacesError("Google refused the key for Places: in Google Cloud, open the key's project -> APIs & "
                              "Services -> Library -> \"Places API (New)\" -> Enable (and, if the key is restricted, "
                              "allow Places API (New) on it). Then try again.") from e
        if e.code == 400 and "API key not valid" in detail:
            raise PlacesError("Google says the API key isn't valid - check GOOGLE_PLACES_API_KEY in Settings.") from e
        if e.code == 429:
            raise PlacesError("Google's Places limit was reached for now - try again later, or raise the quota in Google Cloud.") from e
        raise PlacesError(f"Google Places answered HTTP {e.code}.") from e
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as e:
        raise PlacesError(f"Couldn't reach Google Places ({getattr(e, 'reason', e)}).") from e


def own_site(url: str) -> str:
    """Their website, or "" when it's a directory or social page rather than their own site."""
    from find_prospects import NOT_THEIR_SITE

    url = (url or "").strip()
    return "" if not url or NOT_THEIR_SITE.search(url) else url


def row_for(place: dict, trade_label: str, area: str) -> dict:
    site = own_site(place.get("websiteUri") or "")
    return {
        "Business": ((place.get("displayName") or {}).get("text") or "").strip(),
        "Website": site,
        "Status": "New",
        "Trade": trade_label,
        "Area": area.strip().title(),
        "Phone": (place.get("nationalPhoneNumber") or "").strip(),
        "Registered address": (place.get("formattedAddress") or "").strip(),
        "Website found": "Google Maps" if site else "",
        "Google rating": place.get("rating") or "",
        "Google reviews": place.get("userRatingCount") or 0,
        "Source": "Google Maps",
        "_status": "confirmed" if site else "none",
        "_id": place.get("id") or "",
    }


def keep(place: dict, include: list[str], exclude: list[str]) -> bool:
    name = ((place.get("displayName") or {}).get("text") or "").lower()
    if not name or place.get("businessStatus") not in (None, "OPERATIONAL"):
        return False
    if place.get("primaryType") in NOT_A_PROSPECT or CHAINS.search(name):
        return False
    if include and not any(w in name for w in include):
        return False
    return not any(w and w in name for w in exclude)


def find(trades: list[str], areas: list[str], key: str, known: tuple[set, set, set], most: int,
         include: list[str], exclude: list[str], website_only: bool = False, post=None, log=print) -> list[dict]:
    from find_prospects import TRADES, normalise
    from push_prospects import domain_of

    _, names, domains = known
    seen_ids: set[str] = set()
    rows: list[dict] = []
    for trade in trades:
        for area in areas:
            query = f"{SEARCH_WORDS.get(trade, TRADES[trade]['label'])} in {area}"
            places = search(query, key, post)
            log(f"Google Maps: \"{query}\" - {len(places)} results")
            for p in places:
                if p.get("id") in seen_ids or not keep(p, include, exclude):
                    continue
                seen_ids.add(p.get("id"))
                row = row_for(p, TRADES[trade]["label"], area)
                d = domain_of(row["Website"]) if row["Website"] else None
                if (d and d in domains) or normalise(row["Business"]) in names:
                    continue  # already in one of your lists
                if website_only and not row["Website"]:
                    continue
                rows.append(row)
                if d:
                    domains.add(d)
                names.add(normalise(row["Business"]))
    # Busiest first: plenty of reviews means an established firm that gets work and can pay.
    rows.sort(key=lambda r: -int(r["Google reviews"] or 0))
    return rows[:most]


def write(path: Path, rows: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    from find_prospects import HEADERS, TABS

    headers = HEADERS + EXTRA_HEADERS
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for status, tab in TABS.items():
        if status == "possible":
            continue  # Google gives the website itself - nothing to check
        ws = wb.create_sheet(tab)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "B2"
        for r in rows:
            if r["_status"] == status:
                ws.append([r.get(h, "") for h in headers])
        widths = {"Business": 34, "Website": 30, "Registered address": 44}
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(h, 14)
    wb.save(path)


def run(args, outreach: Path, key: str, post=None, log=print) -> Path | None:
    from find_prospects import known_firms, split_list

    if not key:
        raise PlacesError("Add GOOGLE_PLACES_API_KEY in Settings (or a PAGESPEED_API_KEY from a project with "
                          "Places API (New) enabled).")
    trades = [t.strip() for t in args.trades.split(",") if t.strip()]
    areas = split_list(args.areas)
    include = [w.lower() for w in split_list(args.include)]
    exclude = [w.lower() for w in split_list(args.exclude)]
    rows = find(trades, areas, key, known_firms(outreach), 10_000 if args.count_only else args.max, include, exclude,
                args.website_only, post, log)
    with_site = sum(1 for r in rows if r["Website"])
    if args.count_only:
        log(f"{len(rows)} new firms on Google Maps ({with_site} with a website, {len(rows) - with_site} without).")
        return None
    if not rows:
        log("No new firms found - every match is already in one of your lists.")
        return None
    from find_prospects import output_path

    path = output_path(outreach, ["maps", *trades], areas)
    write(path, rows)
    log(f"Wrote {path.name}: {len(rows)} firms ({with_site} with a website on the Outreach tab, "
        f"{len(rows) - with_site} without on the No website tab), busiest first.")
    log("Next: pick it at the top and press Run the whole list - it finds their emails and checks every site.")
    return path
