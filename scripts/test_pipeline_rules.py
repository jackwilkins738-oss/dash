"""Offline tests for do-not-contact, cross-list duplicates, findings and the Mailmeteor columns.

Run: python -m unittest scripts/test_pipeline_rules.py
"""

import csv
import io
import os
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from datetime import date
from pathlib import Path
from unittest import mock

import openpyxl

from contact_rules import Claims, add_to_blocklist, load_blocklist
from findings import all_issues, top_issue


def domain_of(w):
    from push_prospects import domain_of as d

    return d(w)


def make_sheet(path: Path, rows: list[dict], tab: str = "Outreach") -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab
    header = ["Business", "Website", "Status", "Email", "Company type", "Trade", "Area", "Contact name"]
    for r in rows:
        header += [k for k in r if k not in header]
    ws.append(header)
    for r in rows:
        ws.append([r.get(h, "") for h in header])
    wb.save(path)


class Findings(unittest.TestCase):
    def test_order_matches_the_preview_page(self):
        t = {"checks": {"whatsapp": False, "contactForm": False, "tapToCall": False}, "copyrightYear": 2019}
        issues = all_issues(t, 2026)
        self.assertEqual(issues[0], "your phone number isn't tap-to-call on a mobile")
        self.assertEqual(issues[1], "there's no enquiry form on your homepage")
        self.assertIn("your website footer still says © 2019", issues)
        self.assertEqual(issues[-1], "there's no WhatsApp link on your site")

    def test_nothing_found(self):
        self.assertEqual(top_issue({"checks": {"tapToCall": True}}, 2026), "")
        self.assertEqual(top_issue(None, 2026), "")

    def test_small_image_savings_and_recent_years_are_not_issues(self):
        self.assertEqual(all_issues({"checks": {}, "imageSavingsKb": 200, "copyrightYear": 2025}, 2026), [])


class Rules(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_no_in_one_sheet_blocks_the_firm_everywhere(self):
        make_sheet(self.dir / "old.xlsx", [{"Business": "Smith Roofing Ltd", "Website": "www.smithroofing.co.uk", "Status": "Lost / not interested", "Email": "info@smithroofing.co.uk"}])
        block = load_blocklist(self.dir, domain_of)
        self.assertTrue(block.why("smithroofing.co.uk", "", ""))
        self.assertTrue(block.why("", "INFO@smithroofing.co.uk", ""))
        self.assertTrue(block.why("", "", "SMITH ROOFING"))  # spelled differently in a new list
        self.assertIsNone(block.why("jones.co.uk", "a@jones.co.uk", "Jones"))

    def test_hand_added_entries(self):
        add_to_blocklist(self.dir, [("website", "acme.co.uk"), ("email", "bob@acme.co.uk")], "asked by phone")
        self.assertEqual(add_to_blocklist(self.dir, [("website", "ACME.co.uk")], "again"), 0)  # no duplicates
        block = load_blocklist(self.dir, domain_of)
        self.assertEqual(block.why("acme.co.uk", "", ""), "asked by phone")

    def test_claims_belong_to_the_first_list_and_lapse(self):
        make_sheet(self.dir / "a.xlsx", [{"Business": "X", "Website": "x.co.uk"}])
        listed = {}
        load_blocklist(self.dir, domain_of, listed)
        claims = Claims(self.dir, listed)
        claims.claim(["x.co.uk"], "a.xlsx")
        self.assertEqual(Claims(self.dir, listed).other_owner("x.co.uk", "b.xlsx"), "a.xlsx")
        self.assertIsNone(Claims(self.dir, listed).other_owner("x.co.uk", "a.xlsx"))
        # Removed from a.xlsx: no longer a's.
        make_sheet(self.dir / "a.xlsx", [])
        listed = {}
        load_blocklist(self.dir, domain_of, listed)
        self.assertIsNone(Claims(self.dir, listed).other_owner("x.co.uk", "b.xlsx"))


class PushRun(unittest.TestCase):
    """push_prospects end to end (dry run) with the rules in place."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        make_sheet(self.dir / "master.xlsx", [
            {"Business": "Said No Roofing", "Website": "saidno.co.uk", "Status": "Lost / not interested", "Email": "a@saidno.co.uk", "Company type": "Ltd"},
            {"Business": "Shared Roofing", "Website": "shared.co.uk", "Status": "New", "Email": "a@shared.co.uk", "Company type": "Ltd"},
        ])
        make_sheet(self.dir / "new-list.xlsx", [
            {"Business": "SAID NO ROOFING LTD", "Website": "saidno-roofing.co.uk", "Status": "New", "Email": "x@saidno-roofing.co.uk", "Company type": "Ltd"},
            {"Business": "Shared Roofing", "Website": "https://www.shared.co.uk/", "Status": "New", "Email": "a@shared.co.uk", "Company type": "Ltd"},
            {"Business": "Fresh Roofing", "Website": "fresh.co.uk", "Status": "New", "Email": "info@fresh.co.uk", "Company type": "Ltd", "Trade": "Roofing", "Area": "Woking", "Contact name": "Ann Lee"},
        ])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
            w.writerow(["fresh.co.uk", "2026-09-01", "ok", "41", "6.2", "your phone number isn't tap-to-call on a mobile", "3"])

    def tearDown(self):
        self.tmp.cleanup()

    def run_push(self, sheet):
        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / sheet), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        return buf.getvalue()

    def test_blocked_and_duplicate_firms_are_left_out_of_the_new_list(self):
        self.run_push("master.xlsx")  # the master lists (and so owns) Shared Roofing first
        out = self.run_push("new-list.xlsx")
        self.assertIn("1 prospects", out)
        self.assertIn("they said no", out)
        self.assertIn("SAID NO ROOFING LTD", out)
        self.assertIn("already in master.xlsx", out)
        with (self.dir / "mailmeteor-new-list.csv").open(encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual([r["business"] for r in rows], ["Fresh Roofing"])
        r = rows[0]
        self.assertEqual((r["trade"], r["area"], r["greeting_name"]), ("roofing", "Woking", "Ann"))
        self.assertEqual(r["top_issue"], "your phone number isn't tap-to-call on a mobile")
        self.assertEqual(r["mobile_score"], "41")
        self.assertEqual(r["score_line"], "It scored 41 out of 100 on Google's mobile speed test, which Google itself counts as poor.")
        self.assertEqual(r["issue_line"], "I also noticed your phone number isn't tap-to-call on a mobile.")
    def test_mailmeteor_only_links_to_previews_that_exist(self):
        from push_prospects import make_slug

        live_slug = make_slug("Fresh Roofing", "fresh.co.uk", "x" * 40)
        make_sheet(self.dir / "two.xlsx", [
            {"Business": "Fresh Roofing", "Website": "fresh.co.uk", "Status": "New", "Email": "info@fresh.co.uk", "Company type": "Ltd"},
            {"Business": "Unpushed Roofing", "Website": "unpushed.co.uk", "Status": "New", "Email": "a@unpushed.co.uk", "Company type": "Ltd"},
        ])
        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {live_slug: {}}), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "two.xlsx"), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        with (self.dir / "mailmeteor-two.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["business"] for r in csv.DictReader(f)], ["Fresh Roofing"])
        self.assertIn("1 email firms left out of mailmeteor-two.csv", buf.getvalue())

    def test_the_mailmeteor_file_never_waits_on_first_lines(self):
        from push_prospects import make_slug

        slug = make_slug("Fresh Roofing", "fresh.co.uk", "x" * 40)
        make_sheet(self.dir / "fl.xlsx", [
            {"Business": "Fresh Roofing", "Website": "fresh.co.uk", "Status": "New", "Email": "info@fresh.co.uk", "Company type": "Ltd"},
        ])

        def broken(*a, **k):
            raise RuntimeError("homepage parser fell over")

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40, "ANTHROPIC_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {slug: {}}), \
                mock.patch("first_line.fill", broken), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "fl.xlsx"), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        with (self.dir / "mailmeteor-fl.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["business"] for r in csv.DictReader(f)], ["Fresh Roofing"])
        self.assertIn("First lines skipped this time", buf.getvalue())

    def test_new_first_lines_land_in_the_mailmeteor_file(self):
        from push_prospects import make_slug

        slug = make_slug("Fresh Roofing", "fresh.co.uk", "x" * 40)
        make_sheet(self.dir / "fl2.xlsx", [
            {"Business": "Fresh Roofing", "Website": "fresh.co.uk", "Status": "New", "Email": "info@fresh.co.uk", "Company type": "Ltd"},
        ])
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(io.StringIO()), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {slug: {}}), \
                mock.patch("first_line.fill", lambda *a, **k: {"fresh.co.uk": "Saw you fit flat roofs in Woking."}), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "fl2.xlsx"), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        with (self.dir / "mailmeteor-fl2.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["first_line"] for r in csv.DictReader(f)], ["Saw you fit flat roofs in Woking."])

    def test_a_push_adds_its_firms_to_mailmeteor(self):
        import json as _json

        make_sheet(self.dir / "push.xlsx", [
            {"Business": "New Roofing", "Website": "newroof.co.uk", "Status": "New", "Email": "info@newroof.co.uk", "Company type": "Ltd"},
        ])

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def fake_urlopen(req, timeout=0):
            return Res(_json.dumps({"upserted": 1, "rejected": []}).encode())

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {}), \
                mock.patch("urllib.request.urlopen", fake_urlopen), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "push.xlsx")]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertIn("1 email firms left out", out)  # before the push
        with (self.dir / "mailmeteor-push.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["business"] for r in csv.DictReader(f)], ["New Roofing"])  # after it

    def test_verify_links_keeps_only_links_that_load(self):
        import urllib.error

        from push_prospects import make_slug

        make_sheet(self.dir / "v.xlsx", [
            {"Business": "Good Roofing", "Website": "good.co.uk", "Status": "New", "Email": "a@good.co.uk", "Company type": "Ltd"},
            {"Business": "Broken Roofing", "Website": "broken.co.uk", "Status": "New", "Email": "a@broken.co.uk", "Company type": "Ltd"},
        ])
        slugs = {make_slug(b, w, "x" * 40): {} for b, w in (("Good Roofing", "good.co.uk"), ("Broken Roofing", "broken.co.uk"))}
        seen = []

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def fake_urlopen(req, timeout=0):
            seen.append(req.full_url)
            if "broken" in req.full_url:
                raise urllib.error.HTTPError(req.full_url, 404, "nf", {}, io.BytesIO(b""))
            return Res(b"<title>Prepared for Good Roofing</title>")

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: slugs), \
                mock.patch("push_prospects.urllib.request.urlopen", fake_urlopen), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "v.xlsx"), "--dry-run", "--verify-links"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertIn("Broken Roofing: HTTP 404", out)
        self.assertIn("READY: 1 links checked and loading", out)
        with (self.dir / "mailmeteor-v.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["business"] for r in csv.DictReader(f)], ["Good Roofing"])
        # Checked as the owner's own visit, never counted as theirs.
        self.assertTrue(seen and all("src=dashboard" in u and "src=email" not in u for u in seen))

    def test_link_loads_needs_their_preview(self):
        from push_prospects import link_loads

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        with mock.patch("push_prospects.urllib.request.urlopen", lambda req, timeout=0: Res(b"<title>Scalar Digital</title>")):
            self.assertEqual(link_loads("https://x/for/a?src=email"), "page loaded but isn't their preview")

    def test_companies_house_down_never_stops_the_push(self):
        import json as _json

        from company_lookup import LookupFailed

        make_sheet(self.dir / "ch.xlsx", [
            {"Business": f"Blank Type {i}", "Website": f"blanktype{i}.co.uk", "Status": "New", "Email": f"a@blanktype{i}.co.uk"} for i in range(4)
        ])

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def down(*a, **k):
            raise LookupFailed("HTTP 503")

        buf = io.StringIO()
        env = {"PROSPECTS_API_SECRET": "x" * 40, "COMPANIES_HOUSE_API_KEY": "1a2b3c4d-5e6f-7a8b-9c0d-1e2f3a4b5c6d"}
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), mock.patch("company_lookup.lookup", down), \
                mock.patch("site_teardown.fetch_html", lambda url, timeout=20: (None, None)), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {}), \
                mock.patch("urllib.request.urlopen", lambda req, timeout=0: Res(_json.dumps({"upserted": 1, "rejected": []}).encode())), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "ch.xlsx"), "--lookup-companies"]):
            import push_prospects

            push_prospects.main()  # no SystemExit
        out = buf.getvalue()
        self.assertIn("WARNING: company lookup stopped", out)
        self.assertIn("Pushed 1", out)  # the fake dashboard always answers "1"

    def test_previews_with_only_a_name_are_flagged(self):
        from push_prospects import make_slug

        make_sheet(self.dir / "nm.xlsx", [
            {"Business": "Scored", "Website": "scored.co.uk", "Status": "New", "Email": "a@scored.co.uk", "Company type": "Ltd"},
            {"Business": "Name Only", "Website": "nameonly.co.uk", "Status": "New", "Email": "a@nameonly.co.uk", "Company type": "Ltd"},
        ])
        slugs = {make_slug(b, w, "x" * 40): {} for b, w in (("Scored", "scored.co.uk"), ("Name Only", "nameonly.co.uk"))}

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def page(req, timeout=0):
            if "scored" in req.full_url:
                return Res(b"<title>Prepared for Scored</title> 41 / 100 on mobile ... What we found")
            return Res(b"<title>Prepared for Name Only</title>")

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": ""}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: slugs), \
                mock.patch("push_prospects.urllib.request.urlopen", page), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "nm.xlsx"), "--dry-run", "--verify-links"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertIn("NOTE: 1 of 2 previews load but show no speed score", out)
        self.assertIn("no PAGESPEED_API_KEY", out)
        self.assertIn("    Name Only", out)
        self.assertIn("READY: all 2 links checked", out)  # they still load, so they stay in

    def test_diagnosis_names_the_website_settings(self):
        import urllib.error

        from push_prospects import make_slug

        make_sheet(self.dir / "dg.xlsx", [{"Business": "On Dash", "Website": "ondash.co.uk", "Status": "New", "Email": "a@ondash.co.uk", "Company type": "Ltd"}])
        slug = make_slug("On Dash", "ondash.co.uk", "x" * 40)

        def not_found(req, timeout=0):
            raise urllib.error.HTTPError(req.full_url, 404, "nf", {}, io.BytesIO(b""))

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, secret, tenant: {slug: {}}), \
                mock.patch("push_prospects.urllib.request.urlopen", not_found), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "dg.xlsx"), "--dry-run", "--verify-links"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertIn("DIAGNOSIS: 1 of 1 links don't load.", out)
        self.assertIn("on your dashboard but the website says 'not found'", out)
        self.assertIn("PROSPECTS_API_SECRET", out)

    def test_parallel_speed_checks_record_findings(self):
        import threading

        running, peak = [0], [0]
        lock = threading.Lock()

        def fake_teardown(domain, key):
            with lock:
                running[0] += 1
                peak[0] = max(peak[0], running[0])
            time.sleep(0.05)
            with lock:
                running[0] -= 1
            if domain == "broken.co.uk":
                raise RuntimeError("boom")
            return {"v": 1, "checks": {"contactForm": False}, "_mobile_score": 55, "_lcp_s": 4.1}

        rows = [{"Business": f"Firm {i}", "Website": f"firm{i}.co.uk", "Status": "New", "Company type": "Ltd"} for i in range(7)]
        rows.append({"Business": "Broken", "Website": "broken.co.uk", "Status": "New", "Company type": "Ltd"})
        make_sheet(self.dir / "batch.xlsx", rows)
        buf = io.StringIO()
        env = {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": "k"}
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), mock.patch("site_teardown.teardown", fake_teardown), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "batch.xlsx"), "--teardown", "--yes-all", "--dry-run"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertEqual(peak[0], 4, out)  # four at a time, never more
        self.assertIn("55/100, 1 issue - worst: there's no enquiry form on your homepage", out)
        self.assertIn("broken.co.uk: check failed (RuntimeError)", out)
    def test_a_preview_page_without_its_speed_check_gets_one(self):
        """Checked before under another spelling: the website is in the log, but this list's page never got
        the check. The dashboard says so, and the site is checked again - keeping its last score meanwhile."""
        from push_prospects import make_slug

        make_sheet(self.dir / "leads.xlsx", [
            {"Business": "KERR ROOFING LIMITED", "Website": "kerr.co.uk", "Status": "New", "Company type": "Ltd"},
            {"Business": "Done Roofing", "Website": "done.co.uk", "Status": "New", "Company type": "Ltd"},
        ])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
            w.writerow(["kerr.co.uk", "2026-09-01", "ok", "38", "6.2", "", "0"])
            w.writerow(["done.co.uk", "2026-09-01", "ok", "71", "2.9", "", "0"])
        secret = "x" * 40
        done = make_slug("Done Roofing", "done.co.uk", secret)
        activity = {done: {"teardown_at": "2026-09-01T10:00:00Z"}}  # Kerr's page under this spelling has none
        checked = []

        def fake_teardown(domain, key):
            checked.append(domain)
            return {"v": 1, "checks": {}, "_mobile_score": 40, "_lcp_s": 5.0}

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": secret, "PAGESPEED_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, s, tenant: activity), \
                mock.patch("site_teardown.teardown", fake_teardown), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "leads.xlsx"), "--teardown", "--limit", "5", "--dry-run"]):
            import push_prospects

            push_prospects.main()
        self.assertEqual(checked, ["kerr.co.uk"], buf.getvalue())
        self.assertIn("1 preview page(s) have no speed score on them yet", buf.getvalue())
        self.assertIn("kerr.co.uk: 40/100", buf.getvalue())  # the fresh score, not the remembered 38
        with (self.dir / "teardown-log.csv").open(encoding="utf-8") as f:
            log = {r["website"]: r for r in csv.DictReader(f)}
        self.assertEqual((log["kerr.co.uk"]["result"], log["kerr.co.uk"]["mobile_score"]), ("recheck", "38"))  # dry run: not logged yet
        self.assertEqual(log["done.co.uk"]["result"], "ok")

    def test_google_giving_no_score_is_said_and_tried_again(self):
        make_sheet(self.dir / "ns.xlsx", [{"Business": "Slow Roofing", "Website": "slow.co.uk", "Status": "New", "Company type": "Ltd"},
                                         {"Business": "Old Roofing", "Website": "old.co.uk", "Status": "New", "Company type": "Ltd"}])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
            w.writerow(["old.co.uk", "2026-09-01", "ok", "", "", "", "0"])  # an earlier silent failure
        checked = []

        def fake_teardown(domain, key):
            checked.append(domain)
            return {"v": 1, "checks": {"tapToCall": False}, "_psi_error": "Google's speed test quota is used up for today"}

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        import json as _json

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, s, tenant: {}), \
                mock.patch("site_teardown.teardown", fake_teardown), \
                mock.patch("google_extras.add_to", lambda *a, **k: None), \
                mock.patch("urllib.request.urlopen", lambda req, timeout=0: Res(_json.dumps({"upserted": 2, "rejected": []}).encode())), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "ns.xlsx"), "--teardown", "--limit", "5"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertEqual(sorted(checked), ["old.co.uk", "slow.co.uk"], out)  # the old silent one is retried too
        self.assertIn("no score (Google's speed test quota is used up for today)", out)
        self.assertIn("WARNING: 2 site(s) got no speed score - Google's speed test quota is used up for today", out)
        with (self.dir / "teardown-log.csv").open(encoding="utf-8") as f:
            log = {r["website"]: r["result"] for r in csv.DictReader(f)}
        self.assertEqual(log, {"old.co.uk": "failed", "slow.co.uk": "failed"})  # Retry failed picks them up

    def test_no_email_goes_out_before_the_speed_check(self):
        from push_prospects import make_slug

        secret = "x" * 40
        make_sheet(self.dir / "wait.xlsx", [
            {"Business": "Checked Roofing", "Website": "checked.co.uk", "Status": "New", "Email": "a@checked.co.uk", "Company type": "Ltd"},
            {"Business": "Waiting Roofing", "Website": "waiting.co.uk", "Status": "New", "Email": "a@waiting.co.uk", "Company type": "Ltd"},
        ])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
            w.writerow(["checked.co.uk", "2026-10-01", "ok", "41", "5.0", "", "0"])
        live = {make_slug("Checked Roofing", "checked.co.uk", secret): {}, make_slug("Waiting Roofing", "waiting.co.uk", secret): {}}
        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": secret, "PAGESPEED_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, s, tenant: live), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "wait.xlsx"), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        with (self.dir / "mailmeteor-wait.csv").open(encoding="utf-8") as f:
            self.assertEqual([r["business"] for r in csv.DictReader(f)], ["Checked Roofing"])
        self.assertIn("1 email firms left out of mailmeteor-wait.csv until their speed check has run", buf.getvalue())

    def test_a_google_extras_crash_never_loses_the_speed_checks(self):
        import json as _json

        make_sheet(self.dir / "gx.xlsx", [{"Business": "Churchill Roofing", "Website": "churchill.co.uk", "Status": "New", "Company type": "Ltd"}])

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def boom(*a, **k):
            raise KeyError("displayName")

        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, s, tenant: {}), \
                mock.patch("site_teardown.teardown", lambda dom, key: {"v": 1, "checks": {}, "_mobile_score": 33}), \
                mock.patch("google_extras.add_to", boom), \
                mock.patch("urllib.request.urlopen", lambda req, timeout=0: Res(_json.dumps({"upserted": 1, "rejected": []}).encode())), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "gx.xlsx"), "--teardown", "--limit", "10"]):
            import push_prospects

            push_prospects.main()
        self.assertIn("Google reviews and competitor scores skipped this batch (KeyError", buf.getvalue())
        with (self.dir / "teardown-log.csv").open(encoding="utf-8") as f:
            log = {r["website"]: (r["result"], r["mobile_score"]) for r in csv.DictReader(f)}
        self.assertEqual(log["churchill.co.uk"], ("ok", "33"))

    def test_an_older_dashboard_changes_nothing(self):
        make_sheet(self.dir / "old.xlsx", [{"Business": "Kerr Roofing", "Website": "kerr.co.uk", "Status": "New", "Company type": "Ltd"}])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
            w.writerow(["kerr.co.uk", "2026-09-01", "ok", "38", "6.2", "", "0"])
        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": "k"}), redirect_stdout(buf), \
                mock.patch("calls.fetch_activity", lambda api, s, tenant: {"some-page-abc123": {"status": "new"}}), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "old.xlsx"), "--teardown", "--limit", "5", "--dry-run"]):
            import push_prospects

            push_prospects.main()
        self.assertIn("Nothing left to speed check on this list.", buf.getvalue())

    def test_speed_check_with_nothing_left_is_not_a_failed_step(self):
        make_sheet(self.dir / "done.xlsx", [{"Business": "Fresh Roofing", "Website": "fresh.co.uk", "Status": "New", "Company type": "Ltd"}])
        buf = io.StringIO()
        env = {"PROSPECTS_API_SECRET": "x" * 40, "PAGESPEED_API_KEY": "k"}
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "done.xlsx"), "--teardown", "--limit", "5", "--dry-run"]):
            import push_prospects

            push_prospects.main()  # returns - no SystemExit, so the panel doesn't flag the step
        self.assertIn("Nothing left to speed check on this list.", buf.getvalue())

    def test_find_contacts_fills_blanks_and_never_overrides_the_sheet(self):
        make_sheet(self.dir / "c.xlsx", [
            {"Business": "Blank Roofing", "Website": "blank.co.uk", "Status": "New", "Company type": "Ltd"},
            {"Business": "Has Email Roofing", "Website": "hasemail.co.uk", "Status": "New", "Email": "mine@hasemail.co.uk", "Company type": "Ltd", "Contact name": "Sue Bell"},
        ])
        with (self.dir / "company-lookups.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "result", "number", "company_type"])
            w.writerow(["blank.co.uk", "company", "01234567", "Ltd"])
        pages = {
            "https://blank.co.uk/": ('<a href="tel:01483 222333">x</a><a href="/contact">c</a>', "https://blank.co.uk/"),
            "https://blank.co.uk/contact": ("<p>office@blank.co.uk</p>", "https://blank.co.uk/contact"),
            "https://hasemail.co.uk/": ("<p>other@hasemail.co.uk 01483 999000</p>", "https://hasemail.co.uk/"),
        }
        officers = {"items": [{"name": "BLANK, Tom", "officer_role": "director", "appointed_on": "2010-01-01"}]}
        buf = io.StringIO()
        env = {"PROSPECTS_API_SECRET": "x" * 40, "COMPANIES_HOUSE_API_KEY": "1a2b3c4d-5e6f-7a8b-9c0d-1e2f3a4b5c6d"}
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), \
                mock.patch("site_teardown.fetch_html", lambda url, timeout=20: pages.get(url, (None, None))), \
                mock.patch("company_lookup._get", lambda path, key: officers), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "c.xlsx"), "--find-contacts", "--dry-run"]):
            import push_prospects

            push_prospects.main()
        out = buf.getvalue()
        self.assertIn("Blank Roofing: office@blank.co.uk · 01483 222333 · Tom Blank", out)
        with (self.dir / "mailmeteor-c.csv").open(encoding="utf-8") as f:
            rows = {r["business"]: r for r in csv.DictReader(f)}
        self.assertEqual((rows["Blank Roofing"]["email"], rows["Blank Roofing"]["greeting_name"]), ("office@blank.co.uk", "Tom"))
        # The sheet's own email and contact win over anything found.
        self.assertEqual((rows["Has Email Roofing"]["email"], rows["Has Email Roofing"]["greeting_name"]), ("mine@hasemail.co.uk", "Sue"))


class Letters(unittest.TestCase):
    def test_fill_leaves_out_unknown_facts(self):
        import letters

        body = letters.fill("Dear {greeting},\n\nA. {score_sentence} {issue_sentence} B. {unknown}", {"greeting": "Tom"})
        self.assertEqual(body, "Dear Tom,\n\nA. B. {unknown}")
        self.assertEqual(letters.score_sentence(38), "It scored 38 out of 100 on mobile, which Google itself counts as poor.")
        self.assertEqual(letters.score_sentence(None), "")

    def test_letters_run(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Outreach"
            head = ["Business", "Website", "Status", "Email", "Company type", "Contact name", "Address", "Trade", "Company number"]
            ws.append(head)
            ws.append(["Sole Roofer", "soleroofer.co.uk", "New", "", "Sole trader", "Bill Kerr", "4 Mill Lane, Woking, GU21 1AA", "Roofing", ""])
            ws.append(["No Address Ltd", "noaddress.co.uk", "New", "", "Ltd", "", "", "Roofing", "07654321"])
            ws.append(["Posted Already", "posted.co.uk", "New", "", "Sole trader", "", "1 Road, Town, GU1 1AA", "", ""])
            ws.append(["Email Firm", "emailfirm.co.uk", "New", "info@emailfirm.co.uk", "Ltd", "", "2 Road, Town", "", ""])
            nw = wb.create_sheet("No website")
            nw.append(["Business", "Status", "Company type", "Contact name", "Registered address", "Trade", "Company number"])
            nw.append(["Siteless Builders", "New", "Ltd", "Ann Hale", "9 High St, Guildford, GU1 2BB", "Building & extensions", "01111111"])
            wb.save(d / "letters.xlsx")
            with (d / "company-lookups.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["website", "result", "number", "company_type"])
                w.writerow(["noaddress.co.uk", "company", "07654321", "Ltd"])
            with (d / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s", "top_issue", "issue_count"])
                w.writerow(["soleroofer.co.uk", "2026-09-01", "ok", "34", "7.1", "your phone number isn't tap-to-call on a mobile", "2"])
            import overrides

            with (d / "letters-sent.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["key", "business", "sheet", "posted"])
                w.writerow([overrides.row_key({"Business": "Posted Already"}), "Posted Already", "letters.xlsx", "2026-09-01"])
            office = {"registered_office_address": {"address_line_1": "Unit 3", "locality": "Leeds", "postal_code": "LS1 1AA"}}
            buf = io.StringIO()
            env = {"PROSPECTS_API_SECRET": "x" * 40, "COMPANIES_HOUSE_API_KEY": "1a2b3c4d-5e6f-7a8b-9c0d-1e2f3a4b5c6d"}
            with mock.patch.dict(os.environ, env), redirect_stdout(buf), mock.patch("company_lookup._get", lambda path, key: office), \
                    mock.patch.object(sys, "argv", ["p", "--sheet", str(d / "letters.xlsx"), "--letters", "--dry-run"]):
                import push_prospects

                push_prospects.main()
            out = buf.getvalue()
            self.assertIn("Wrote letters-letters.html: 3 letters ready", out)
            self.assertIn("1 left out - already posted", out)
            page = (d / "letters-letters.html").read_text(encoding="utf-8")
            self.assertIn("Dear Bill,", page)
            self.assertIn("It scored 34 out of 100 on mobile, which Google itself counts as poor.", page)
            self.assertIn("I also noticed your phone number isn&#x27;t tap-to-call on a mobile.", page)
            self.assertIn("Unit 3<br>Leeds<br>LS1 1AA", page)  # registered office from Companies House
            self.assertIn("Dear Ann,", page)
            self.assertIn("I couldn&#x27;t find a website for Siteless Builders", page)
            self.assertIn("builders", page)
            self.assertNotIn("Email Firm", page)  # goes by email, not letter
            self.assertNotIn("Posted Already", page)
            self.assertEqual(page.count("data:image/svg+xml"), 3)
            with (d / "letters-batch-letters.csv").open(encoding="utf-8") as f:
                self.assertEqual(len(list(csv.DictReader(f))), 3)


class CallList(unittest.TestCase):
    def test_joins_views_with_local_phones_and_letter_dates(self):
        import calls
        import overrides
        from datetime import datetime, timezone
        from push_prospects import make_slug

        secret = "x" * 40
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Outreach"
            ws.append(["Business", "Website", "Status", "Phone", "Contact name"])
            ws.append(["Hot Roofing", "hot.co.uk", "New", "01483 111111", "Amy Hot"])
            ws.append(["Warm Roofing", "warm.co.uk", "New", "01483 222222", ""])
            ws.append(["Quiet Roofing", "quiet.co.uk", "New", "", ""])
            ws.append(["Called Roofing", "called.co.uk", "New", "01483 444444", ""])
            ws.append(["No Roofing", "no.co.uk", "New", "01483 555555", ""])
            wb.save(d / "s.xlsx")
            with (d / "contacts-found.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["website", "email", "phone", "contact", "address", "checked_at"])
                w.writerow(["quiet.co.uk", "", "07700 900000", "", "", ""])
            with (d / "letters-sent.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["key", "business", "sheet", "posted"])
                w.writerow([overrides.row_key({"Business": "Quiet Roofing"}), "Quiet Roofing", "s.xlsx", "2026-09-10"])
            calls.log_call(d, "s.xlsx", overrides.row_key({"Business": "No Roofing"}), "No Roofing", "Not interested")
            now = datetime(2026, 9, 27, 12, tzinfo=timezone.utc)
            act = {
                make_slug("Hot Roofing", "hot.co.uk", secret): {"view_count": 3, "last_viewed_at": "2026-09-27T11:00:00Z"},
                make_slug("Warm Roofing", "warm.co.uk", secret): {"view_count": 1, "last_viewed_at": "2026-09-20T11:00:00Z"},
                make_slug("No Roofing", "no.co.uk", secret): {"view_count": 5, "last_viewed_at": "2026-09-27T11:00:00Z"},
            }
            out = calls.call_list(d, "s.xlsx", secret, "https://site", act, now)
            self.assertEqual([i["business"] for i in out["viewing"]], ["Hot Roofing", "Warm Roofing"])
            hot = out["viewing"][0]
            self.assertEqual((hot["phone"], hot["contact"], hot["views"]), ("01483 111111", "Amy Hot", 3))
            self.assertTrue(hot["preview"].endswith("?src=dashboard"))  # your own look isn't counted as theirs
            self.assertEqual([(i["business"], i["phone"], i["waited"]) for i in out["letters"]], [("Quiet Roofing", "07700 900000", 17)])
            # Called after their last visit: off the list until they look again.
            with (d / "calls.csv").open("a", newline="", encoding="utf-8") as f:
                csv.writer(f).writerow([overrides.row_key({"Business": "Hot Roofing"}), "s.xlsx", "Hot Roofing", "No answer", "", "2026-09-27T12:30:00+00:00"])
            later = datetime(2026, 9, 27, 13, tzinfo=timezone.utc)
            self.assertEqual([i["business"] for i in calls.call_list(d, "s.xlsx", secret, "https://site", act, later)["viewing"]], ["Warm Roofing"])

    def test_dashboard_without_the_endpoint(self):
        import calls
        import urllib.error

        err = urllib.error.HTTPError("u", 404, "nf", {}, io.BytesIO(b""))
        with mock.patch.object(calls.urllib.request, "urlopen", side_effect=err):
            with self.assertRaisesRegex(calls.DashboardMissing, "activity update"):
                calls.fetch_activity("https://x", "s", "t")


class ExportResults(unittest.TestCase):
    def test_dated_groups_notes_kept_and_one_file(self):
        import export_results
        from datetime import date

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            make_sheet(d / "roofing-guildford.xlsx", [
                {"Business": "Kerr Roofing", "Website": "kerr.co.uk", "Status": "New", "Email": "info@kerr.co.uk", "Company type": "Ltd", "Contact name": "Bill Kerr"},
                {"Business": "Sole Roofer", "Website": "sole.co.uk", "Status": "New", "Company type": "Sole trader"},
            ])
            env = {"PROSPECTS_API_SECRET": "x" * 40, "DASHBOARD_API_URL": "http://127.0.0.1:9"}
            with mock.patch.dict(os.environ, env), mock.patch.object(export_results, "outreach_dir", lambda: d), redirect_stdout(io.StringIO()):
                export_results.main()
                # A later list, and a note typed into the workbook in between.
                make_sheet(d / "lofts-woking.xlsx", [{"Business": "Loft Co", "Website": "loftco.co.uk", "Status": "New", "Company type": "Ltd", "Email": "a@loftco.co.uk"}])
                wb = openpyxl.load_workbook(d / "outreach-results.xlsx")
                ws = wb["Results"]
                header = [c.value for c in ws[1]]
                for row in ws.iter_rows(min_row=2):
                    if row[header.index("Business")].value == "Kerr Roofing":
                        row[header.index("Notes")].value = "Rang - call back Tuesday"
                wb.save(d / "outreach-results.xlsx")
                # The first list was exported "earlier": its firms keep that date.
                idx = (d / "results-index.csv").read_text(encoding="utf-8").replace(date.today().isoformat(), "2026-09-20")
                (d / "results-index.csv").write_text(idx, encoding="utf-8")
                export_results.main()

            wb = openpyxl.load_workbook(d / "outreach-results.xlsx")
            ws = wb["Results"]
            header = [c.value for c in ws[1]]
            values = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
            headings = [r[0] for r in values if isinstance(r[0], str) and "·" in r[0]]
            self.assertEqual(len(headings), 2)
            self.assertIn("1 firm", headings[0])  # newest group first: today's lofts list
            self.assertIn("lofts-woking", headings[0])
            self.assertTrue(headings[1].startswith("Sunday 20 September 2026"))
            self.assertIn("2 firms", headings[1])
            self.assertIn(None, [r[0] for r in values])  # a blank spacer row between date groups
            kerr = next(r for r in values if len(r) > 2 and r[header.index("Business")] == "Kerr Roofing")
            self.assertEqual(kerr[header.index("Notes")], "Rang - call back Tuesday")
            self.assertEqual(kerr[header.index("Contact")], "Bill Kerr")
            self.assertEqual(kerr[header.index("Channel")], "Email")
            self.assertEqual(kerr[header.index("Added")].date(), date(2026, 9, 20))
            key_col = openpyxl.utils.get_column_letter(header.index("Key") + 1)
            self.assertTrue(ws.column_dimensions[key_col].hidden)
            summary = [[c.value for c in r] for r in wb["Summary"].iter_rows()]
            self.assertEqual(sum(1 for r in summary if r[0] == "Export"), 2)
            self.assertIn(["All lists", 3, 2, 1], [r[:4] for r in summary])


class EmailBatches(unittest.TestCase):
    def test_batches_take_the_next_firms_and_never_repeat(self):
        import email_batches as eb
        from datetime import date

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            head = ["business", "greeting_name", "email", "preview_url"]
            with (d / "mailmeteor-old.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(head)
                for i in range(15):
                    w.writerow([f"Old {i}", "there", f"a{i}@old.co.uk", f"https://s/for/old-{i}?src=email"])
            with (d / "mailmeteor-new.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(head)
                for i in range(15):
                    w.writerow([f"New {i}", "there", f"b{i}@new.co.uk", f"https://s/for/new-{i}?src=email"])
            import os as _os
            _os.utime(d / "mailmeteor-old.csv", (1, 1))  # the older list goes first
            broken = lambda url: "HTTP 404" if url.endswith("old-3?src=email") else None  # noqa: E731

            path, notes = eb.make_batch(d, 20, check=broken, today=date(2026, 9, 28))
            names = [r["business"] for r in csv.DictReader(path.open(encoding="utf-8"))]
            self.assertEqual(path.name, "mailmeteor-batch-2026-09-28.csv")
            self.assertEqual(len(names), 20)
            self.assertEqual(names[:3], ["Old 0", "Old 1", "Old 2"])
            self.assertNotIn("Old 3", names)
            self.assertIn("skipped Old 3: link HTTP 404", notes)
            self.assertEqual(eb.mark_sent(d, date(2026, 9, 28)), 20)

            path2, _ = eb.make_batch(d, 20, check=lambda u: None, today=date(2026, 9, 29))
            names2 = [r["business"] for r in csv.DictReader(path2.open(encoding="utf-8"))]
            self.assertEqual(set(names) & set(names2), set())  # nobody emailed twice
            self.assertIn("Old 3", names2)  # its link loads now, so it gets its turn
            self.assertEqual(len(names2), 10)
            eb.mark_sent(d, date(2026, 9, 29))
            self.assertEqual(eb.remaining(d), 0)
            self.assertEqual(eb.make_batch(d, 20, check=lambda u: None)[0], None)

    def test_one_list_only(self):
        import email_batches as eb

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            for name, who in (("mailmeteor-a.csv", "A"), ("mailmeteor-b.csv", "B")):
                with (d / name).open("w", newline="", encoding="utf-8") as f:
                    w = csv.writer(f)
                    w.writerow(["business", "email", "preview_url"])
                    w.writerow([who, f"x@{who.lower()}.co.uk", "https://s/for/x"])
            path, _ = eb.make_batch(d, 20, sheet="b.xlsx", check=lambda u: None)
            self.assertEqual([r["business"] for r in csv.DictReader(path.open(encoding="utf-8"))], ["B"])


class FollowUps(unittest.TestCase):
    def test_who_gets_one_and_never_a_third(self):
        import email_batches as eb
        from datetime import date

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            with (d / "emails-sent.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(eb.SENT_FIELDS)
                for who, sent in (("quiet", "2026-09-20"), ("replied", "2026-09-20"), ("viewer", "2026-09-20"),
                                  ("recent", "2026-09-26"), ("blocked", "2026-09-20"), ("ooo", "2026-09-20")):
                    w.writerow([f"{who}@x.co.uk", who.title(), f"https://s/for/{who}-abc123?src=email", "b1", sent, ""])
            with (d / "mailmeteor-l.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["business", "greeting_name", "email", "issue_line", "preview_url"])
                w.writerow(["Quiet", "Sam", "quiet@x.co.uk", "I also noticed your footer still says © 2019.", "https://s/for/quiet-abc123?src=email"])
            with (d / "replies.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["message_id", "date", "from", "business", "website", "kind", "subject", "snippet", "handled"])
                w.writerow(["1", "2026-09-22", "replied@x.co.uk", "Replied", "", "read it", "", "", ""])
                w.writerow(["2", "2026-09-22", "ooo@x.co.uk", "Ooo", "", "out of office", "", "", ""])
            add_to_blocklist(d, [("email", "blocked@x.co.uk")], "said no")
            views = {"viewer-abc123": {"view_count": 2}}
            path, notes, skipped = eb.make_followups(d, 20, 5, check=lambda u: None, views=views, today=date(2026, 9, 28))
            rows = list(csv.DictReader(path.open(encoding="utf-8")))
            self.assertEqual(sorted(r["email"] for r in rows), ["ooo@x.co.uk", "quiet@x.co.uk"])  # an auto-reply isn't a reply
            quiet = next(r for r in rows if r["email"] == "quiet@x.co.uk")
            self.assertEqual((quiet["greeting_name"], quiet["issue_line"]), ("Sam", "I also noticed your footer still says © 2019."))
            self.assertEqual(skipped, {"replied": 1, "viewed": 1, "blocked": 1, "too soon": 1})
            self.assertEqual(eb.mark_followups_sent(d, date(2026, 9, 28)), 2)
            later, _, _ = eb.make_followups(d, 20, 5, check=lambda u: None, views=views, today=date(2026, 10, 30))
            with later.open(encoding="utf-8") as f:
                # The one that was too recent, and the viewer nobody rang in 10 days - never a third email.
                self.assertEqual(sorted(r["email"] for r in csv.DictReader(f)), ["recent@x.co.uk", "viewer@x.co.uk"])
            # A later first-email batch keeps the follow-up dates.
            with (d / eb.PENDING).open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["email", "business", "preview_url", "batch"])
                w.writerow(["new@x.co.uk", "New", "https://s/for/new", "b2"])
            eb.mark_sent(d, date(2026, 9, 29))
            sent = {r["email"]: r for r in csv.DictReader((d / "emails-sent.csv").open(encoding="utf-8"))}
            self.assertEqual(sent["quiet@x.co.uk"]["followup_sent"], "2026-09-28")


class Autopilot(unittest.TestCase):
    def test_a_lock_left_by_a_stopped_or_crashed_run_is_not_running(self):
        import subprocess

        import autopilot
        import control_panel as panel

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            lock = d / autopilot.LOCK
            with mock.patch.object(panel, "OUTREACH", d):
                self.assertFalse(autopilot.is_running())
                # A live Python process holds it: running.
                lock.write_text(str(os.getpid()))
                self.assertTrue(autopilot.is_running())
                # The run was killed (Stop, a shutdown): its process is gone, so the lock is cleared.
                dead = subprocess.Popen([sys.executable, "-c", "pass"])
                dead.wait()
                lock.write_text(str(dead.pid))
                self.assertFalse(autopilot.is_running())
                self.assertFalse(lock.exists())
                # A lock with no process number (an old version's crash) - cleared too.
                lock.write_text("")
                self.assertFalse(autopilot.is_running())
                # Older than 6 hours: cleared even if the number now belongs to something else.
                lock.write_text(str(os.getpid()))
                old = time.time() - 7 * 3600
                os.utime(lock, (old, old))
                self.assertFalse(autopilot.is_running())

    def test_runs_every_step_in_order_and_carries_on(self):
        self._run_autopilot(monday=False)

    def test_on_mondays_the_scorecard_goes_in_the_text(self):
        order, log = self._run_autopilot(monday=True)
        self.assertEqual(order[-1], "scorecard.py")
        self.assertIn("Weekly scorecard:", log)
        self.assertIn("All time: 40 sent", log)

    def _run_autopilot(self, monday):
        import autopilot
        import control_panel as panel

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "panel.env").write_text("PROSPECTS_API_SECRET=" + "x" * 40 + "\nCOMPANIES_HOUSE_API_KEY=1a2b3c4d-5e6f-7a8b-9c0d-1e2f3a4b5c6d\n"
                                         "MAIL_ADDRESS=me@x.co.uk\nMAIL_APP_PASSWORD=abcd\n", encoding="utf-8")
            autopilot.save_config.__globals__["panel"].OUTREACH  # noqa: B018 - same module object
            ran = []

            def fake_step(self, script, args):
                ran.append((script.name, [a for a in args if a.startswith("--")]))
                self.lines.append("READY: mailmeteor-batch-2026-09-28.csv - 20 firms" if script.name == "email_batches.py" and "--followups" not in args else "")
                if script.name == "find_prospects.py":
                    make_sheet(d / "roofing-guildford-2026-09-28.xlsx", [{"Business": "A", "Website": "a.co.uk", "Status": "New"}])
                if script.name == "reply_scanner.py":
                    self.failed.append("reply_scanner.py")  # a failed step: the rest still run
                    return 1
                if script.name == "scorecard.py":
                    self.lines.append("  All time: 40 sent · 3% replied (1)")
                return 0

            with mock.patch.object(panel, "OUTREACH", d), mock.patch.object(panel, "SETTINGS_FILE", d / "panel.env"), \
                    mock.patch.object(autopilot.Run, "step", fake_step), mock.patch.object(panel, "keep_awake", lambda on: None), \
                    mock.patch.object(autopilot, "scorecard_day", lambda: monday), redirect_stdout(io.StringIO()):
                autopilot.save_config({"find": {"trades": ["roofing"], "areas": "Guildford", "age": "any", "max": 50, "exclude": ""},
                                       "batch_size": 20, "followups": True})
                code = autopilot.run()
            order = [name for name, _ in ran]
            self.assertEqual(order[0], "reply_scanner.py")
            self.assertEqual(order[1], "sending_health.py")  # checked before anything can be sent
            self.assertEqual(order[2], "find_prospects.py")
            self.assertIn("push_prospects.py", order)  # the whole list ran on the new sheet
            tail = order[:-1] if monday else order
            self.assertEqual(tail[-4:], ["email_batches.py", "email_batches.py", "export_results.py", "backup.py"])
            self.assertIn(["--followups", "--size", "--after-days"], [a for n, a in ran if n == "email_batches.py"])
            self.assertEqual(code, 1)  # one step had a problem, and it says so
            self.assertFalse((d / autopilot.LOCK).exists())
            log = (d / "autopilot-log.txt").read_text(encoding="utf-8")
            self.assertIn("READY: mailmeteor-batch", log)
            self.assertIn("1 step(s) had problems", log)
            return order, log

    def test_the_windows_schedule_command(self):
        import autopilot

        seen = {}

        def fake_run(cmd, capture_output=True, text=True):
            seen["cmd"] = cmd
            return mock.MagicMock(returncode=0, stdout="SUCCESS", stderr="")

        with mock.patch.object(autopilot.sys, "platform", "win32"), mock.patch.object(autopilot.subprocess, "run", fake_run):
            self.assertEqual(autopilot.install("06:45"), "")
            self.assertEqual(autopilot.install("6.45"), "The time must look like 07:30.")
        cmd = seen["cmd"]
        self.assertEqual(cmd[:4], ["schtasks", "/Create", "/F", "/SC"])
        self.assertEqual(cmd[cmd.index("/ST") + 1], "06:45")
        self.assertEqual(cmd[cmd.index("/TN") + 1], "Scalar Prospect Autopilot")
        self.assertIn("autopilot.py", cmd[cmd.index("/TR") + 1])

    def test_settings_are_checked_before_saving(self):
        import control_panel as panel

        with tempfile.TemporaryDirectory() as tmp:
            with mock.patch.object(panel, "OUTREACH", Path(tmp)):
                bad, code = panel.autopilot_action({"action": "save", "config": {"time": "7.30"}})
                self.assertEqual(code, 400)
                bad, code = panel.autopilot_action({"action": "save", "config": {"time": "07:30", "find": {"trades": [], "areas": "x"}}})
                self.assertIn("Search", bad["error"])
                ok, code = panel.autopilot_action({"action": "save", "config": {"time": "06:45", "batch_size": 25, "followups": True,
                                                   "find": {"trades": ["roofing"], "areas": "Woking", "max": 40}}})
                self.assertEqual(code, 200)
                import autopilot

                self.assertEqual(autopilot.load_config()["find"]["areas"], "Woking")


class Scorecard(unittest.TestCase):
    def test_versions_trades_and_a_fair_verdict(self):
        import csv as _csv

        import scorecard

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            def write(name, fields, rows):
                with (d / name).open("w", newline="", encoding="utf-8") as f:
                    w = _csv.DictWriter(f, fieldnames=fields)
                    w.writeheader()
                    w.writerows(rows)
            sent = [{"email": f"a{i}@x.co.uk", "business": f"Firm {i}", "preview_url": f"https://s/for/firm-{i}",
                     "sent": "2026-09-25", "variant": "A" if i % 2 else "B"} for i in range(240)]
            write("emails-sent.csv", ["email", "business", "preview_url", "sent", "variant"], sent)
            write("mailmeteor.csv", ["email", "trade"], [{"email": r["email"], "trade": "Roofing" if i < 120 else "Loft conversions"}
                                                         for i, r in enumerate(sent)])
            # B gets 8 replies, A gets 1; one bounce and one out-of-office don't count.
            replies = [{"from": f"a{i}@x.co.uk", "business": f"Firm {i}", "kind": "interested" if i < 4 else "read it"}
                       for i in (0, 2, 4, 6, 8, 10, 12, 14)]
            replies += [{"from": "a1@x.co.uk", "business": "Firm 1", "kind": "not interested"},
                        {"from": "a3@x.co.uk", "business": "Firm 3", "kind": "bounce"},
                        {"from": "a5@x.co.uk", "business": "Firm 5", "kind": "out of office"}]
            write("replies.csv", ["from", "business", "kind"], replies)
            write("quotes-sent.csv", ["business"], [{"business": "Firm 0"}])
            write("calls.csv", ["business", "outcome", "at"], [{"business": "Firm 0", "outcome": "Won", "at": "2026-09-27T10:00:00+00:00"},
                                                                 {"business": "Firm 7", "outcome": "No answer", "at": "2026-09-27T11:00:00+00:00"}])
            views = {f"firm-{i}": {"view_count": 1} for i in range(0, 240, 4)}
            lines = scorecard.scorecard(d, views, date(2026, 9, 28))
        text = "\n".join(lines)
        self.assertIn("All time: 240 sent · 25% opened · 4% replied (9) · 2 keen · 1 quoted · 1 won", text)
        self.assertIn("Version B: 120 sent", text)
        self.assertIn("replied (8)", text)
        self.assertIn("Version B is getting clearly more replies", text)
        self.assertIn("Roofing: 120 sent", text)
        self.assertIn("Calls in the last 7 days: 2, real conversations: 1", text)

    def test_too_early_to_call_is_said_plainly(self):
        import scorecard

        rows = [{"variant": v, "replied": i < 3} for i, v in enumerate("AB" * 30)]
        self.assertIn("Too early", scorecard.verdict(rows))

    def test_versions_are_split_about_evenly_and_stick_to_an_address(self):
        import send_email

        t = {**send_email.DEFAULT_TEMPLATES, "first_subject_b": "Your site, {{business}}", "first_body_b": "Hi {{greeting_name}} {{preview_url}}"}
        picks = [send_email.variant_for(f"firm{i}@x.co.uk", t) for i in range(1000)]
        self.assertTrue(400 < picks.count("B") < 600)
        self.assertEqual(send_email.variant_for("Firm1@X.co.uk", t), send_email.variant_for("firm1@x.co.uk", t))
        self.assertEqual(send_email.variant_for("firm1@x.co.uk", send_email.DEFAULT_TEMPLATES), "A")
        self.assertIn("both a subject and a body", send_email.template_problem({**send_email.DEFAULT_TEMPLATES, "first_subject_b": "x"}))
        self.assertIn("preview_url", send_email.template_problem({**t, "first_body_b": "Hi {{business}}"}))
        self.assertEqual(send_email.template_problem(t), "")


class Quotes(unittest.TestCase):
    def test_email_it_to_them_sends_the_recorded_link_from_your_email(self):
        import control_panel as panel
        import quotes

        sent = []

        class FakeSMTP:
            def send_message(self, msg):
                sent.append(msg)

            def quit(self):
                pass

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            make_sheet(d / "list.xlsx", [{"Business": "Kerr Roofing", "Website": "kerr.co.uk"}])
            (d / "panel.env").write_text("MAIL_ADDRESS=jack@scalar.co.uk\nMAIL_APP_PASSWORD=abcd\nMAIL_FROM_NAME=Jack Wilkins\n", encoding="utf-8")
            quotes.record(d, "list.xlsx", "name:kerr roofing", "Kerr Roofing", "build",
                          {"quote_number": "SD-0007", "total_pence": 250000, "quote_url": "https://admin.scalardigital.co.uk/quote/q1/t1"})
            with mock.patch.object(panel, "OUTREACH", d), mock.patch.object(panel, "SETTINGS_FILE", d / "panel.env"), \
                 mock.patch("send_email.connect", return_value=FakeSMTP()):
                ok, code = panel.quote_email_action({"sheet": "list.xlsx", "key": "name:kerr roofing", "to": "bill@kerr.co.uk",
                                                     "contact": "Bill Kerr", "quote_number": "SD-0007"})
                missing, code2 = panel.quote_email_action({"sheet": "list.xlsx", "key": "name:kerr roofing", "to": "bill@kerr.co.uk",
                                                           "quote_number": "SD-9999"})
        self.assertEqual(code, 200, ok)
        msg = sent[0]
        self.assertEqual((msg["To"], msg["From"]), ("bill@kerr.co.uk", "Jack Wilkins <jack@scalar.co.uk>"))
        self.assertEqual(msg["Subject"], "Your website quote - Kerr Roofing (SD-0007)")
        body = msg.get_content()
        self.assertIn("Hi Bill,", body)
        self.assertIn("https://admin.scalardigital.co.uk/quote/q1/t1", body)
        self.assertIn("(£2,500)", body)
        self.assertIn("Jack Wilkins", body)
        self.assertEqual(code2, 404)
        self.assertEqual(len(sent), 1)

    def test_what_is_sent_and_the_email_it_opens(self):
        import json as _json
        import urllib.parse

        import quotes

        sent = {}

        class Res(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def fake_urlopen(req, timeout=0):
            sent["url"], sent["body"], sent["auth"] = req.full_url, _json.loads(req.data), req.headers.get("Authorization")
            return Res(_json.dumps({"ok": True, "quote_number": "Q-0007", "total_pence": 250000,
                                    "quote_url": "https://admin.scalardigital.co.uk/quote/q1/t1"}).encode())

        settings = {"PROSPECTS_API_SECRET": "x" * 40, "QUOTE_PRICE_BUILD": "2500", "QUOTE_VAT_RATE": "0", "QUOTE_DEPOSIT_PERCENT": "50"}
        with mock.patch.object(quotes.urllib.request, "urlopen", fake_urlopen):
            out = quotes.create_quote(settings, "Kerr Roofing", "info@kerr.co.uk", "01483 111222", "build", "kerr-roofing-4a7bc2")
        self.assertEqual(sent["url"], "https://admin.scalardigital.co.uk/api/prospects/quote")
        self.assertEqual(sent["auth"], "Bearer " + "x" * 40)
        b = sent["body"]
        self.assertEqual((b["client_name"], b["slug"], b["vat_rate"], b["deposit_percent"]), ("Kerr Roofing", "kerr-roofing-4a7bc2", 0, 50))
        self.assertEqual(b["line_items"][0]["unit_price_pence"], 250000)
        self.assertIn("A 50% deposit is due when you accept", b["terms"])
        self.assertTrue(b["exclusions"] and b["payment_terms"].startswith("50% deposit"))
        self.assertEqual(b["optional_items"][0], {"description": "Extra page", "price_pence": 15000})
        link = quotes.email_link("info@kerr.co.uk", "Bill Kerr", "Kerr Roofing", out, "Jack")
        self.assertTrue(link.startswith("mailto:info%40kerr.co.uk?subject="))
        body = urllib.parse.unquote(link.split("body=", 1)[1])
        self.assertIn("Hi Bill,", body)
        self.assertIn("(£2,500)", body)
        self.assertIn("https://admin.scalardigital.co.uk/quote/q1/t1", body)

    def test_quote_extras_from_settings(self):
        import quotes

        self.assertEqual(len(quotes.extras({})), 4)  # blank: the defaults
        self.assertEqual(quotes.extras({"QUOTE_EXTRAS": "none"}), [])
        self.assertEqual(quotes.extras({"QUOTE_EXTRAS": "Extra page £175; Logo 99.50; no price here; Free 0"}),
                         [{"description": "Extra page", "price_pence": 17500}, {"description": "Logo", "price_pence": 9950}])

    def test_a_dashboard_without_the_update_says_so(self):
        import urllib.error

        import quotes

        err = urllib.error.HTTPError("u", 404, "nf", {}, io.BytesIO(b""))
        with mock.patch.object(quotes.urllib.request, "urlopen", side_effect=err):
            with self.assertRaisesRegex(quotes.QuoteFailed, "doesn't have the quote update"):
                quotes.create_quote({"PROSPECTS_API_SECRET": "x" * 40}, "A", "", "", "landing")

    def test_panel_quote_logs_the_call_and_clears_the_reply(self):
        import control_panel as panel
        import quotes

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            make_sheet(d / "s.xlsx", [{"Business": "Kerr Roofing", "Website": "kerr.co.uk", "Status": "New"}])
            (d / "panel.env").write_text("PROSPECTS_API_SECRET=" + "x" * 40 + "\n", encoding="utf-8")
            with (d / "replies.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["message_id", "date", "from", "business", "website", "kind", "subject", "snippet", "handled"])
                w.writerow(["1", "2026-09-28", "bill@kerr.co.uk", "Kerr Roofing", "kerr.co.uk", "interested", "", "", ""])
            fake = {"quote_number": "Q-0001", "total_pence": 75000, "quote_url": "https://admin.x/quote/1/t"}
            with mock.patch.object(panel, "OUTREACH", d), mock.patch.object(panel, "SETTINGS_FILE", d / "panel.env"), \
                    mock.patch.object(quotes, "create_quote", lambda *a, **k: fake):
                out, code = panel.quote_action({"sheet": "s.xlsx", "key": "name:kerr roofing", "business": "Kerr Roofing",
                                                "email": "bill@kerr.co.uk", "website": "kerr.co.uk", "package": "landing"})
                bad, bad_code = panel.quote_action({"sheet": "s.xlsx", "key": "name:kerr roofing", "business": "Kerr Roofing", "package": "gold"})
            self.assertEqual((code, out["quote_number"], out["url"]), (200, "Q-0001", "https://admin.x/quote/1/t"))
            self.assertEqual(bad_code, 400)
            with (d / "calls.csv").open(encoding="utf-8") as f:
                self.assertEqual([r["outcome"] for r in csv.DictReader(f)], ["Quoted"])
            with (d / "replies.csv").open(encoding="utf-8") as f:
                self.assertTrue(next(csv.DictReader(f))["handled"].startswith("answered"))
            with (d / "quotes-sent.csv").open(encoding="utf-8") as f:
                self.assertEqual(next(csv.DictReader(f))["total"], "750.00")


class SiteDraft(unittest.TestCase):
    PAGE = """<html><head><title>Kerr Roofing | Roofers in Guildford</title>
<meta name="description" content="Family roofers since 1998.">
<meta name="theme-color" content="#b3261e"><style>.x{color:#ffffff}</style></head><body>
<nav><a href="/">Home</a><a href="/flat-roofs">Flat Roofs</a><a href="/repairs">Roof Repairs</a><a href="/contact">Contact us</a>
<a href="/chimneys">Chimney &amp; Leadwork</a></nav>
<img src="/img/kerr-logo.png" alt="Kerr logo"><img src="/img/job1.jpg"><img src="/icons/phone.png">
<h1>Roofing in Guildford <script>x</script></h1><h2>Why choose us</h2>
<p>Established in 1998. NFRC members.</p><img src="/img/job2.webp"></body></html>"""

    def test_reads_their_site(self):
        import site_draft

        s = site_draft.read_site(self.PAGE, "https://kerr.co.uk/")
        self.assertEqual(s["services"], ["Flat Roofs", "Roof Repairs", "Chimney & Leadwork"])
        self.assertEqual(s["colour"], "#b3261e")
        self.assertEqual(s["logo"], "https://kerr.co.uk/img/kerr-logo.png")
        self.assertEqual(s["photos"], ["https://kerr.co.uk/img/job1.jpg", "https://kerr.co.uk/img/job2.webp"])
        self.assertEqual((s["years"], s["accreditations"]), ("since 1998", ["NFRC"]))

    def test_preview_only_gets_what_we_are_sure_of(self):
        import site_draft

        sure = {"services": ["Flat Roofs", "Roof Repairs", "A very long service name that won't fit", "Chimneys"],
                "colour": "#b3261e", "colour_sure": True}
        self.assertEqual(site_draft.for_preview(sure), {"services": ["Flat Roofs", "Roof Repairs", "Chimneys"], "brandColour": "#b3261e"})
        with_images = {"logo": "https://kerr.co.uk/img/kerr-logo.png",
                       "photos": ["https://kerr.co.uk/img/job1.jpg", "http://kerr.co.uk/img/job2.jpg", "https://kerr.co.uk/checkatrade-member.png",
                                  "https://kerr.co.uk/img/job3.webp", "https://kerr.co.uk/img/plan.svg", "https://kerr.co.uk/a.jpg",
                                  "https://kerr.co.uk/b.jpg", "https://kerr.co.uk/c.jpg"]}
        out = site_draft.for_preview(with_images)
        self.assertEqual(out["logo"], "https://kerr.co.uk/img/kerr-logo.png")
        self.assertEqual(out["photos"], ["https://kerr.co.uk/img/job1.jpg", "https://kerr.co.uk/img/job3.webp",
                                         "https://kerr.co.uk/a.jpg", "https://kerr.co.uk/b.jpg"])
        self.assertNotIn("logo", site_draft.for_preview({"logo": "https://kerr.co.uk/logo.svg"}))
        self.assertEqual(site_draft.for_preview({"services": ["Flat Roofs"], "colour": "#b3261e", "colour_sure": False}), {})
        # A colour used once in their CSS isn't a brand colour; one too pale for white text is darkened, same hue.
        once = site_draft.read_site('<style>.a{color:#2a7ab0}</style><a href="/r">Roof repairs</a>', "https://x.co.uk/")
        self.assertFalse(once["colour_sure"])
        self.assertEqual(site_draft.readable("#ffcc00"), "#866a00")
        self.assertGreaterEqual(site_draft._contrast_with_white(site_draft.readable("#87ceeb")), 4.5)

    def test_trade_keys(self):
        import site_draft

        self.assertEqual([site_draft.trade_key(t) for t in ["Roofing", "Loft conversions", "Driveways & patios", "Landscaping",
                                                            "Building & extensions", "Plumbing", ""]],
                         ["roofing", "lofts", "driveways", "landscaping", "building", "general", "general"])

    def test_writes_a_brief_and_an_escaped_draft(self):
        import site_draft

        with tempfile.TemporaryDirectory() as d:
            out = Path(d)
            make_sheet(out / "list.xlsx", [{"Business": "Kerr <b>Roofing</b>", "Website": "kerr.co.uk", "Trade": "Roofing",
                                            "Area": "Guildford", "Phone": "01483 111222", "Email": "info@kerr.co.uk",
                                            "Company number": "1234567", "Registered name": "KERR ROOFING LTD",
                                            "Registered address": "1 High St, Guildford", "Incorporated": "2010-04-01"}])
            with mock.patch("site_teardown.fetch_html", return_value=(self.PAGE, "https://kerr.co.uk/")):
                folder = site_draft.make(out, "list.xlsx", "no:01234567")
            page = (folder / "index.html").read_text(encoding="utf-8")
            brief = (folder / "brief.md").read_text(encoding="utf-8")
        self.assertEqual(folder.name, "kerr-b-roofing-b")
        self.assertNotIn("<b>Roofing</b>", page)
        self.assertIn("Kerr &lt;b&gt;Roofing&lt;/b&gt;", page)
        self.assertIn('href="tel:01483111222"', page)
        self.assertIn("--brand: #b3261e", page)
        self.assertIn("<h3>Flat Roofs</h3>", page)
        self.assertIn("company no. 1234567", page)
        self.assertIn('name="robots" content="noindex"', page)
        self.assertIn("**Services they list:** Flat Roofs, Roof Repairs", brief)
        self.assertIn("https://kerr.co.uk/img/job1.jpg", brief)

    def test_what_they_sent_on_the_onboarding_page_replaces_the_guesses(self):
        import site_draft

        told = {"services": "Flat roofs\nRe-roofing\n- Chimneys", "areas": "Guildford\nWoking\nGodalming", "phone": "07700 900123",
                "guarantee": "10 years on new roofs", "insurance": "Yes", "insurance_amount": "£5 million",
                "memberships": "NFRC, TrustMark", "about": "Family run since 1998."}
        files = [{"kind": "logo", "name": "logo.png", "path": "client-files/logo-01-logo.png"},
                 {"kind": "photo", "name": "a.jpg", "path": "client-files/photo-02-a.jpg"},
                 {"kind": "photo", "name": "b.heic", "path": "client-files/photo-03-b.heic"}]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d)
            make_sheet(out / "list.xlsx", [{"Business": "Kerr Roofing", "Website": "kerr.co.uk", "Trade": "Roofing", "Phone": "01483 111222"}])
            with mock.patch("site_teardown.fetch_html", return_value=(None, None)), \
                 mock.patch("site_draft.fetch_onboarding", return_value={"answers": told, "submitted_at": "2026-10-01T09:00:00Z", "files": []}) as fo, \
                 mock.patch("site_draft.download_files", return_value=files):
                folder = site_draft.make(out, "list.xlsx", "name:kerr roofing", settings={"PROSPECTS_API_SECRET": "s" * 40})
            page = (folder / "index.html").read_text(encoding="utf-8")
            brief = (folder / "brief.md").read_text(encoding="utf-8")
        self.assertRegex(fo.call_args.args[1], r"^kerr-roofing-[0-9a-f]{6}$")
        self.assertIn("<h3>Chimneys</h3>", page)
        self.assertIn("Serving Guildford, Woking and Godalming", page)
        self.assertIn('href="tel:07700900123"', page)
        self.assertIn("<b>Guarantee: 10 years on new roofs</b>", page)
        self.assertIn("(£5 million public liability)", page)
        self.assertIn("<li><b>TrustMark</b></li>", page)
        self.assertIn('src="client-files/logo-01-logo.png"', page)
        self.assertIn('src="client-files/photo-02-a.jpg"', page)
        self.assertNotIn("photo-03-b.heic", page)  # browsers can't show HEIC; it's kept in the folder
        self.assertNotIn("Ask for 6-12 photos", page)
        self.assertIn("About Kerr Roofing", page)
        self.assertIn("<p>Yes - 10 years on new roofs.</p>", page)
        self.assertIn("What they told us (onboarding page, sent 2026-10-01)", brief)

    def test_no_site_and_no_row_falls_back(self):
        import site_draft

        with tempfile.TemporaryDirectory() as d:
            out = Path(d)
            make_sheet(out / "list.xlsx", [{"Business": "Other", "Website": "other.co.uk"}])
            with mock.patch("site_teardown.fetch_html", return_value=(None, None)):
                folder = site_draft.make(out, "list.xlsx", "name:acme drives",
                                         {"business": "Acme Drives", "website": "acme.co.uk", "phone": ""})
            page = (folder / "index.html").read_text(encoding="utf-8")
            brief = (folder / "brief.md").read_text(encoding="utf-8")
            with self.assertRaises(ValueError):
                site_draft.make(out, "list.xlsx", "name:nobody", {})
        self.assertIn("Phone number</mark>", page)
        self.assertIn("[Their main service]".strip("[]"), page)
        self.assertIn("Couldn't read their homepage", brief)


class Review(unittest.TestCase):
    def setUp(self):
        import openpyxl

        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outreach"
        head = ["Business", "Website", "Status", "Email", "Company type", "Company number"]
        ws.append(head)
        ws.append(["Unsure Roofing", "unsure.co.uk", "New", "a@unsure.co.uk", "", ""])
        ws.append(["Bad Email Roofing", "bademail.co.uk", "New", "bob@gmial.com", "Ltd", ""])
        ws.append(["Down Roofing", "down.co.uk", "New", "", "Ltd", ""])
        ws.append(["Gone Roofing", "gone.co.uk", "New", "", "", ""])
        chk = wb.create_sheet("Check website")
        chk.append(["Business", "Website", "Possible website", "Status", "Email", "Company type", "Company number", "Website found"])
        chk.append(["Maybe Roofing", "", "mayberoofing.co.uk", "New", "info@mayberoofing.co.uk", "Ltd", "01234567", "name only"])
        wb.save(self.dir / "list.xlsx")
        with (self.dir / "company-lookups.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "result", "how", "status", "registered_name", "number"])
            w.writerow(["unsure.co.uk", "unsure", "unsure (1 same-name company, not confirmed)", "", "", ""])
            w.writerow(["gone.co.uk", "closed", "number on their site", "dissolved", "GONE ROOFING LTD", "07777777"])
        with (self.dir / "email-checks.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["email", "result", "checked_at"])
            w.writerow(["bob@gmial.com", "typo (did they mean gmail.com?)", "2026-09-01"])
        with (self.dir / "teardown-log.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["website", "checked_at", "result", "mobile_score", "lcp_s"])
            w.writerow(["down.co.uk", "2026-09-01", "failed", "", ""])

    def tearDown(self):
        self.tmp.cleanup()

    def push(self):
        buf = io.StringIO()
        with mock.patch.dict(os.environ, {"PROSPECTS_API_SECRET": "x" * 40}), redirect_stdout(buf), \
                mock.patch.object(sys, "argv", ["p", "--sheet", str(self.dir / "list.xlsx"), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        with (self.dir / "preview-links-list.csv").open(encoding="utf-8") as f:
            return buf.getvalue(), {r["Business"]: r for r in csv.DictReader(f)}

    def test_finds_every_kind(self):
        import review

        kinds = {(i["kind"], i["business"]) for i in review.items(self.dir, "list.xlsx")}
        self.assertEqual(kinds, {
            ("website", "Maybe Roofing"), ("company", "Unsure Roofing"), ("email", "Bad Email Roofing"),
            ("failed", "Down Roofing"), ("closed", "Gone Roofing"),
        })

    def test_decisions_change_the_next_run_and_leave_the_review(self):
        import review

        items = {i["business"]: i for i in review.items(self.dir, "list.xlsx")}
        _, before = self.push()
        self.assertEqual(before["Unsure Roofing"]["Channel"], "letter")
        self.assertNotIn("Maybe Roofing", before)

        review.decide(self.dir, "list.xlsx", items["Unsure Roofing"]["key"], "set_type", "Ltd")
        review.decide(self.dir, "list.xlsx", items["Maybe Roofing"]["key"], "website_yes", "mayberoofing.co.uk")
        review.decide(self.dir, "list.xlsx", items["Bad Email Roofing"]["key"], "set_email", "bob@gmail.com")
        review.decide(self.dir, "list.xlsx", items["Down Roofing"]["key"], "skip")
        out, after = self.push()
        self.assertEqual(after["Unsure Roofing"]["Channel"], "email")
        self.assertEqual(after["Maybe Roofing"]["Channel"], "email")
        self.assertEqual(after["Bad Email Roofing"]["Email"], "bob@gmail.com")
        self.assertNotIn("Down Roofing", after)
        self.assertIn("1 left out in Review", out)
        left = {i["business"] for i in review.items(self.dir, "list.xlsx")}
        self.assertEqual(left, {"Gone Roofing"})

        review.decide(self.dir, "list.xlsx", items["Down Roofing"]["key"], "undo")
        self.assertIn("Down Roofing", self.push()[1])

    def test_do_not_contact_from_review(self):
        import review

        review.block(self.dir, "Unsure Roofing", "unsure.co.uk", "a@unsure.co.uk")
        out, rows = self.push()
        self.assertNotIn("Unsure Roofing", rows)
        self.assertIn("do not contact (Review)", out)


class Panel(unittest.TestCase):
    def test_phone_alert_carries_counts_never_names(self):
        import control_panel

        sent = {}

        def fake_urlopen(req, timeout=0):
            import json

            sent.update(json.loads(req.data))
            return mock.MagicMock()

        lines = [
            "[3/10] smithroofing.co.uk: 41/100, 3 issues - worst: your phone number ...",
            "    Smith Roofing: info@smithroofing.co.uk",
            "Pushed 10 (rejected: [])",
            "List done: 120 speed checked, 4 couldn't be checked.",
        ]
        with mock.patch.object(control_panel.urllib.request, "urlopen", fake_urlopen):
            control_panel.phone_alert("Run the whole list", 0, False, 3900, lines, {"TELEGRAM_BOT_TOKEN": "t", "TELEGRAM_CHAT_ID": "1"})
        self.assertIn("Done: Run the whole list (1h 5m)", sent["text"])
        self.assertIn("List done: 120", sent["text"])
        self.assertNotIn("smith", sent["text"].lower())

    def test_progress_uses_the_same_rules(self):
        import control_panel

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            make_sheet(d / "master.xlsx", [{"Business": "No", "Website": "no.co.uk", "Status": "Lost / not interested"}])
            make_sheet(d / "new.xlsx", [
                {"Business": "No", "Website": "no.co.uk", "Status": "New"},
                {"Business": "Yes", "Website": "yes.co.uk", "Status": "New"},
            ])
            with mock.patch.object(control_panel, "OUTREACH", d):
                self.assertEqual(control_panel.progress("new.xlsx")["total"], 1)

    def test_whole_list_carries_on_past_a_failed_step(self):
        import control_panel

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "fail.py").write_text("import sys; print('boom'); sys.exit(1)")
            (d / "ok.py").write_text("print('pushed fine')")
            job = control_panel.Job()
            err = job.start("Run the whole list", lambda j: [(d / "fail.py", []), (d / "ok.py", [])], {"PROSPECTS_API_SECRET": "x" * 40}, keep_going=True)
            self.assertEqual(err, "")
            job.thread.join(20)
            lines = "\n".join(job.state()["lines"])
            self.assertIn("pushed fine", lines)
            self.assertIn("carrying on with the rest", lines)
            self.assertIn("1 step(s) had problems", lines)
            # Without keep_going, a failure still stops a single action.
            job2 = control_panel.Job()
            job2.start("x", lambda j: [(d / "fail.py", []), (d / "ok.py", [])], {"PROSPECTS_API_SECRET": "x" * 40})
            job2.thread.join(20)
            self.assertNotIn("pushed fine", "\n".join(job2.state()["lines"]))

    def test_no_alert_without_telegram(self):
        import control_panel

        with mock.patch.object(control_panel.urllib.request, "urlopen") as u:
            control_panel.phone_alert("x", 0, False, 999, ["Pushed 1"], {})
        u.assert_not_called()


if __name__ == "__main__":
    unittest.main()
