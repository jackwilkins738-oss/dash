"""What in a sheet needs a decision from you - the panel's Review tab.

Gathers, for one sheet, everything the automatic steps couldn't settle:
    website   a finder sheet's "Check website" rows: is this site theirs?
    company   Companies House couldn't confirm the type (unsure / no record):
              they're getting a letter - say Ltd if you know they are
    email     the email check says their address can't take mail
    failed    their site couldn't be speed checked
    closed    the register says dissolved / in liquidation: left out
Each item carries the firm's key (overrides.row_key) so a decision sticks
across runs, and links to check it yourself.
"""

from __future__ import annotations

import csv
from pathlib import Path
from urllib.parse import quote_plus

import overrides

KINDS = {
    "website": "Is this their website?",
    "company": "Company type not confirmed",
    "email": "Email can't take mail",
    "failed": "Speed check failed",
    "closed": "Dissolved or in liquidation",
}


def _csv(path: Path, key: str) -> dict[str, dict]:
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        return {r[key]: r for r in csv.DictReader(f) if r.get(key)}


def items(outreach: Path, sheet: str) -> list[dict]:
    import openpyxl

    from email_check import is_bad
    from push_prospects import SKIP_STATUSES, domain_of

    path = outreach / sheet
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def tab(name: str) -> list[dict]:
        if name not in wb.sheetnames:
            return []
        rows = list(wb[name].iter_rows(values_only=True))
        header = [str(h).strip() if h else "" for h in rows[0]] if rows else []
        return [dict(zip(header, v)) for v in rows[1:]]

    outreach_rows, check_rows = tab("Outreach"), tab("Check website")
    wb.close()
    decided = overrides.load(outreach, sheet)
    lookups = _csv(outreach / "company-lookups.csv", "website")
    emails = _csv(outreach / "email-checks.csv", "email")
    log = _csv(outreach / "teardown-log.csv", "website")

    out: list[dict] = []

    def add(kind: str, row: dict, detail: str, **extra) -> None:
        business = str(row.get("Business") or "").strip()
        out.append({
            "kind": kind,
            "key": overrides.row_key(row),
            "business": business,
            "website": domain_of(str(row.get("Website") or "")) or "",
            "email": str(row.get("Email") or "").strip(),
            "phone": str(row.get("Phone") or "").strip(),
            "detail": detail,
            "ch_search": "https://find-and-update.company-information.service.gov.uk/search/companies?q=" + quote_plus(business),
            **extra,
        })

    for row in check_rows:
        d = decided.get(overrides.row_key(row), {})
        if "website" in d or d.get("skip") == "yes":
            continue
        site = domain_of(str(row.get("Possible website") or ""))
        if site:
            add("website", row, str(row.get("Website found") or "their name is on it, nothing local"), possible=site)

    for row in outreach_rows:
        if str(row.get("Status") or "") in SKIP_STATUSES:
            continue
        d = decided.get(overrides.row_key(row), {})
        if d.get("skip") == "yes":
            continue
        row = overrides.apply(row, d) or row
        domain = domain_of(str(row.get("Website") or ""))
        if not domain:
            continue
        if not str(row.get("Company type") or "").strip():
            found = lookups.get(domain) or {}
            if found.get("result") in ("unsure", "none"):
                add("company", row, found.get("how") or found["result"], result=found["result"])
            elif found.get("result") == "closed":
                add("closed", row, f"{found.get('registered_name') or ''} - {found.get('status') or 'closed'}".strip(" -"),
                    number=found.get("number") or "")
        email = str(row.get("Email") or "").strip().lower()
        if email and "email" not in d:
            result = (emails.get(email) or {}).get("result") or ""
            if result and is_bad(result):
                add("email", row, result)
        if (log.get(domain) or {}).get("result") == "failed":
            add("failed", row, "the site didn't load for Google's test or ours")

    order = list(KINDS)
    out.sort(key=lambda i: (order.index(i["kind"]), i.get("result") != "unsure", i["business"].lower()))
    return out


def decide(outreach: Path, sheet: str, key: str, action: str, value: str = "") -> str:
    """Applies one Review button. Returns a short confirmation."""
    if action == "set_type":
        overrides.save(outreach, sheet, key, "company_type", value)
        return f"Company type set to {value}"
    if action == "keep_closed":
        overrides.save(outreach, sheet, key, "company_type", "Ltd")
        return "Kept in - treated as Ltd"
    if action == "website_yes":
        overrides.save(outreach, sheet, key, "website", value)
        return f"{value} confirmed - they join the list"
    if action == "website_no":
        overrides.save(outreach, sheet, key, "skip", "yes")
        return "Left out"
    if action == "set_email":
        overrides.save(outreach, sheet, key, "email", value)
        return f"Email changed to {value}" if value else "Email removed - they'll get a letter"
    if action == "skip":
        overrides.save(outreach, sheet, key, "skip", "yes")
        return "Left out of this list"
    if action == "undo":
        overrides.remove(outreach, sheet, key)
        return "Decisions undone"
    raise ValueError("unknown action")


def block(outreach: Path, business: str, website: str, email: str) -> str:
    from contact_rules import add_to_blocklist

    entries = [("website", website), ("email", email), ("name", business)]
    add_to_blocklist(outreach, [(k, v) for k, v in entries if v], "do not contact (Review)")
    return f"{business} won't be contacted again, from any list"
