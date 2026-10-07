"""Canvass cards: four A6 cards to an A4 page, one per firm, each with a QR code to their own preview.

The panel's "Make canvass cards" (Letters card). For a day out in one town: drop them at yards, vans
and offices - "I've made something for you" in their hand, with the page already built. Picks the
list's firms in that area (blank: all), the slowest sites first, leaving out anyone who said no,
replied or is already a client. Opens as an HTML page: print it on card, cut in four.

Each scan shows as src=card on the preview, so the Calls tab knows a card worked.
"""

from __future__ import annotations

import argparse
import csv
import html
import os
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

MOST = 40


def pick(outreach: Path, sheet: str, area: str, secret: str, most: int = MOST) -> list[dict]:
    import openpyxl

    import calls
    import overrides
    from contact_rules import load_blocklist
    from push_prospects import SKIP_STATUSES, domain_of, make_slug

    scores = {}
    log = outreach / "teardown-log.csv"
    if log.exists():
        with log.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                try:
                    scores[r["website"]] = int(float(r.get("mobile_score") or ""))
                except (KeyError, ValueError):
                    pass
    history = calls.load_calls(outreach)
    block = load_blocklist(outreach, domain_of)
    wb = openpyxl.load_workbook(outreach / sheet, read_only=True, data_only=True)
    try:
        if "Outreach" not in wb.sheetnames:
            return []
        rows = wb["Outreach"].iter_rows(values_only=True)
        header = [str(h).strip() if h else "" for h in next(rows, ())]
        out = []
        for values in rows:
            row = dict(zip(header, values))
            business = str(row.get("Business") or "").strip()
            domain = domain_of(str(row.get("Website") or ""))
            town = str(row.get("Area") or row.get("Town") or "").strip()
            if not business or not domain or str(row.get("Status") or "") in SKIP_STATUSES:
                continue
            if area and area.lower() not in f"{town} {row.get('Address') or ''}".lower():
                continue
            if block.why(domain, str(row.get("Email") or ""), business):
                continue
            last = (history.get(overrides.row_key(row)) or [{}])[-1]
            if last.get("outcome") in calls.FINAL or last.get("outcome") == "Interested":
                continue
            address = str(row.get("Address") or row.get("Registered address") or "").strip()
            out.append({"business": business, "area": town, "score": scores.get(domain), "address": address,
                        "slug": make_slug(business, domain, secret)})
    finally:
        wb.close()
    out.sort(key=lambda r: (r["score"] is None, r["score"] if r["score"] is not None else 0, r["business"].lower()))
    return route(out[:most])


def _postcode(address: str) -> str:
    from company_lookup import POSTCODE

    found = POSTCODE.findall((address or "").upper())
    return "".join(found[-1]) if found else ""


def route(firms: list[dict]) -> list[dict]:
    """The day's cards in driving order: by postcode, so neighbours are next to each other (no address: last)."""
    def key(pair: tuple[int, dict]) -> tuple:
        i, f = pair
        pc = _postcode(f.get("address", ""))
        m = re.match(r"([A-Z]{1,2})(\d+)([A-Z]?)(\d)?", pc)
        return (not pc, (m.group(1), int(m.group(2)), m.group(3), m.group(4) or "") if m else ("",), i)  # no address: as picked

    return [f for _, f in sorted(enumerate(firms), key=key)]


def maps_links(firms: list[dict], per_link: int = 10) -> list[str]:
    """Google Maps directions through the addresses, in card order - one link per 10 stops."""
    from urllib.parse import quote

    stops = [f["address"] for f in firms if _postcode(f.get("address", ""))]
    return ["https://www.google.com/maps/dir/" + "/".join(quote(s, safe="") for s in stops[i:i + per_link])
            for i in range(0, len(stops), per_link)]


def render(firms: list[dict], site: str, sender: dict) -> str:
    import letters

    cards = []
    for f in firms:
        url = f"{site}/for/{f['slug']}?src=card"
        short = url.split("://", 1)[-1].split("?")[0]
        cards.append(f"""<div class="card">
  <p class="for">Prepared for</p>
  <h2>{html.escape(f['business'])}</h2>
  <p class="msg">I've built a preview of a faster website for you. Scan it with your phone's camera - nothing to sign up for.</p>
  <img src="{letters.qr_svg(url)}" alt="QR code to your preview">
  <p class="url">{html.escape(short)}</p>
  <p class="me">{html.escape(sender['name'])} · Scalar Digital · {html.escape(sender['phone'])}</p>
</div>""")
    pages = ["<section class=\"page\">" + "".join(cards[i:i + 4]) + "</section>" for i in range(0, len(cards), 4)]
    links = maps_links(firms)
    route_html = ("<nav class=\"route\">Today's route, in card order: " + " ".join(
        f'<a href="{html.escape(u)}" target="_blank" rel="noopener">Stops {i * 10 + 1}-{min((i + 1) * 10, len(firms))} ↗</a>'
        for i, u in enumerate(links)) + "</nav>") if links else ""
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Canvass cards</title>
<style>
@page {{ size: A4; margin: 0; }}
body {{ margin: 0; font-family: Arial, sans-serif; color: #111; }}
.page {{ width: 210mm; height: 297mm; display: grid; grid-template-columns: 1fr 1fr; grid-template-rows: 1fr 1fr;
  page-break-after: always; }}
.card {{ box-sizing: border-box; padding: 10mm; border: 0.2mm dashed #bbb; display: flex; flex-direction: column;
  align-items: center; text-align: center; }}
.for {{ margin: 0; font-size: 9pt; letter-spacing: .15em; text-transform: uppercase; color: #555; }}
h2 {{ margin: 3mm 0; font-size: 17pt; }}
.msg {{ margin: 0 0 5mm; font-size: 10.5pt; line-height: 1.4; }}
img {{ width: 42mm; height: 42mm; }}
.url {{ margin: 3mm 0 0; font-size: 9pt; color: #333; word-break: break-all; }}
.me {{ margin-top: auto; font-size: 9pt; color: #333; }}
.route {{ padding: 4mm 10mm; font-size: 11pt; }} .route a {{ margin-right: 4mm; }}
@media print {{ .route {{ display: none; }} }}
</style></head><body>{route_html}{''.join(pages)}</body></html>"""


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--area", default="")
    ap.add_argument("--most", type=int, default=MOST)
    ap.add_argument("--outreach", type=Path, default=next(
        (d / "outreach" for d in [HERE.parent, *HERE.parents] if (d / "outreach").is_dir()), HERE.parent / "outreach"))
    args = ap.parse_args(argv)
    try:
        import segno  # noqa: F401
    except ImportError:
        raise SystemExit("Cards need one extra piece: python -m pip install segno  (the panel has an Install button).") from None
    secret = os.environ.get("PROSPECTS_API_SECRET", "")
    if len(secret) < 32:
        raise SystemExit("Add PROSPECTS_API_SECRET in Settings first.")
    firms = pick(args.outreach, args.sheet, args.area.strip(), secret, args.most)
    if not firms:
        raise SystemExit(f"No firms in {args.sheet}{' around ' + args.area if args.area else ''} to make cards for.")
    site = os.environ.get("SITE_URL", "https://www.scalardigital.co.uk").rstrip("/")
    sender = {"name": os.environ.get("LETTER_SIGNOFF", "") or os.environ.get("MAIL_FROM_NAME", "") or "Scalar Digital",
              "phone": os.environ.get("LETTER_PHONE", "") or "07401 696272"}
    out = args.outreach / f"cards-{Path(args.sheet).stem}{'-' + args.area.lower().replace(' ', '-') if args.area else ''}.html"
    out.write_text(render(firms, site, sender), encoding="utf-8")
    print(f"Wrote {out.name}: {len(firms)} card{'s' if len(firms) != 1 else ''} ({(len(firms) + 3) // 4} A4 page(s)) - "
          "open it, print on card, cut in four. Their preview pages must already be live (Run the whole list first).")
    for i, link in enumerate(maps_links(firms)):
        print(f"Route, stops {i * 10 + 1}+ (in card order): {link}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
