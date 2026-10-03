"""places_finder: sole traders from Google Maps, busiest first, chains and known firms left out."""

import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import openpyxl

import places_finder


def place(pid, name, site="", reviews=0, status="OPERATIONAL", ptype="roofing_contractor", phone="01483 111222"):
    return {"id": pid, "displayName": {"text": name}, "websiteUri": site, "nationalPhoneNumber": phone,
            "formattedAddress": "1 High St, Guildford GU1 1AA", "rating": 4.8, "userRatingCount": reviews,
            "businessStatus": status, "primaryType": ptype}


PAGE1 = {"places": [place("a", "Kerr Roofing", "https://kerrroofing.co.uk/", 120),
                    place("b", "Smith Roofs", "https://www.facebook.com/smithroofs", 40),
                    place("c", "Wickes Guildford", "https://wickes.co.uk", 900, ptype="home_improvement_store"),
                    place("d", "Closed Roofing", "https://closed.co.uk", 10, status="CLOSED_PERMANENTLY")],
         "nextPageToken": "t2"}
PAGE2 = {"places": [place("e", "Busy Roofers Ltd", "https://busy.co.uk", 300), place("a", "Kerr Roofing", "https://kerrroofing.co.uk/", 120)]}


class Fake:
    def __init__(self):
        self.bodies = []

    def __call__(self, body, key):
        self.bodies.append(body)
        return PAGE2 if body.get("pageToken") == "t2" else PAGE1


class Finder(unittest.TestCase):
    def test_pages_filters_and_order(self):
        fake = Fake()
        rows = places_finder.find(["roofing"], ["Guildford"], "k", (set(), set(), set()), 100, [], ["holdings"], post=fake, log=lambda *_: None)
        self.assertEqual(fake.bodies[0]["textQuery"], "roofer in Guildford")
        self.assertEqual(fake.bodies[1]["pageToken"], "t2")
        self.assertEqual([r["Business"] for r in rows], ["Busy Roofers Ltd", "Kerr Roofing", "Smith Roofs"])
        smith = rows[2]
        self.assertEqual((smith["Website"], smith["_status"]), ("", "none"))  # a Facebook page isn't their site
        self.assertEqual((rows[1]["Phone"], rows[1]["Trade"], rows[1]["Area"]), ("01483 111222", "Roofing", "Guildford"))

    def test_known_firms_and_website_only(self):
        known = (set(), {"kerr roofing"}, {"busy.co.uk"})
        rows = places_finder.find(["roofing"], ["Guildford"], "k", known, 100, [], [], website_only=True, post=Fake(), log=lambda *_: None)
        self.assertEqual(rows, [])

    def test_max(self):
        rows = places_finder.find(["roofing"], ["Guildford"], "k", (set(), set(), set()), 1, [], [], post=Fake(), log=lambda *_: None)
        self.assertEqual([r["Business"] for r in rows], ["Busy Roofers Ltd"])

    def test_workbook_tabs_and_pipeline(self):
        d = Path(tempfile.mkdtemp())
        args = SimpleNamespace(trades="roofing", areas="Guildford", include="", exclude="", website_only=False, count_only=False, max=100)
        lines = []
        path = places_finder.run(args, d, "k", post=Fake(), log=lines.append)
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Outreach", "No website"])
        out = [r for r in wb["Outreach"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(out), 2)
        self.assertEqual(len(list(wb["No website"].iter_rows(min_row=2))), 1)
        self.assertIn("busiest first", lines[-2])
        # The rest of the pipeline reads it like any list: a dry run makes the Mailmeteor and links files.
        env = {"PROSPECTS_API_SECRET": "x" * 40, "DASHBOARD_API_URL": "http://127.0.0.1:9"}
        import io
        from contextlib import redirect_stdout

        buf = io.StringIO()
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), mock.patch.object(sys, "argv", ["p", "--sheet", str(path), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        self.assertIn("2 prospects", buf.getvalue())  # the one with no website is a letter, not a preview

    def test_no_key(self):
        args = SimpleNamespace(trades="roofing", areas="Guildford", include="", exclude="", website_only=False, count_only=False, max=100)
        with self.assertRaises(places_finder.PlacesError):
            places_finder.run(args, Path(tempfile.mkdtemp()), "")

    def test_count_only_writes_nothing(self):
        d = Path(tempfile.mkdtemp())
        args = SimpleNamespace(trades="roofing", areas="Guildford", include="", exclude="", website_only=False, count_only=True, max=100)
        lines = []
        self.assertIsNone(places_finder.run(args, d, "k", post=Fake(), log=lines.append))
        self.assertIn("3 new firms on Google Maps (2 with a website, 1 without)", lines[-1])
        self.assertEqual(list(d.glob("*.xlsx")), [])


class PanelArgs(unittest.TestCase):
    def test_source_google_needs_a_google_key_not_companies_house(self):
        import control_panel as panel

        body = {"action": "find", "source": "google", "trades": ["roofing"], "areas": "Guildford", "max": 50}
        args, err = panel.build_find_args(body, {"PAGESPEED_API_KEY": "k"})
        self.assertEqual(err, "")
        self.assertEqual(args[-2:], ["--source", "google"])
        args, err = panel.build_find_args(body, {})
        self.assertIn("GOOGLE_PLACES_API_KEY", err)


if __name__ == "__main__":
    unittest.main()
