"""Area scout: which towns are worth a list, measured rather than guessed.

The panel's Find card, "Scout areas". For each town, and each trade ticked:

    Google Maps   one search ("roofer in Guildford") - the firms Google shows first, with their
                  review counts. 10+ reviews = established: it gets work and can pay for a site.
    PageSpeed     Google's mobile score for up to 6 of each town's busiest firms with a website.

Each town gets an opportunity figure: established firms with a website you haven't contacted yet,
times the share of the measured ones scoring under 50 - plus half a point for each established firm
with no website at all (a letter). Ranked highest first, into outreach/area-scout.xlsx, with the top
of the table in the log.

Leave Areas blank for the South East commuter towns below - where homes are worth most, so jobs are
bigger. Searches and scores are kept 30 days, so a second scout is quick and costs nothing.

Needs GOOGLE_PLACES_API_KEY (or PAGESPEED_API_KEY with Places API (New) enabled) and PAGESPEED_API_KEY.
"""

from __future__ import annotations

import argparse
import json
import os
import statistics
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

CACHE = "area-scout-cache.json"
OUT = "area-scout.xlsx"
FRESH_DAYS = 30
ESTABLISHED = 10  # Google reviews
SAMPLE = 6  # speed checks per town
SLOW = 50
DEFAULT_TOWNS = [
    "Guildford", "Woking", "Farnham", "Godalming", "Epsom", "Leatherhead", "Dorking", "Reigate", "Esher",
    "Weybridge", "Camberley", "Horsham", "Sevenoaks", "Tunbridge Wells", "Tonbridge", "Maidstone", "St Albans",
    "Harpenden", "Rickmansworth", "Amersham", "Beaconsfield", "High Wycombe", "Marlow", "Maidenhead", "Windsor",
    "Wokingham", "Reading", "Basingstoke", "Winchester", "Chichester",
]
TRADE_QUERY = {"roofing": "roofer", "lofts": "loft conversion company", "driveways": "driveway contractor",
               "landscaping": "landscaper", "building": "builder"}
COLUMNS = ["Rank", "Town", "Opportunity", "Firms on Google Maps", "Established (10+ reviews)", "With a website",
           "Speed checked", "Median score", "Under 50", "Established, no website", "Already in your lists",
           "Busiest firm's reviews"]


def _fresh(at: str) -> bool:
    try:
        return (date.today() - date.fromisoformat(at[:10])).days < FRESH_DAYS
    except (TypeError, ValueError):
        return False


def load_cache(outreach: Path) -> dict:
    try:
        data = json.loads((outreach / CACHE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    return {"searches": data.get("searches") or {}, "scores": data.get("scores") or {}}


def save_cache(outreach: Path, cache: dict) -> None:
    tmp = outreach / (CACHE + ".tmp")
    tmp.write_text(json.dumps(cache), encoding="utf-8")
    tmp.replace(outreach / CACHE)


def firms_in(town: str, trades: list[str], key: str, cache: dict, post=None) -> list[dict]:
    """Every firm Google Maps shows first for these trades in this town, once each."""
    from places_finder import keep, own_site, search
    from push_prospects import domain_of

    seen, firms = set(), []
    for trade in trades:
        query = f"{TRADE_QUERY.get(trade, trade)} in {town}"
        hit = cache["searches"].get(query)
        if not (hit and _fresh(hit.get("at", ""))):
            places = [p for p in search(query, key, post, pages=1) if keep(p, [], [])]
            hit = {"at": date.today().isoformat(), "firms": [
                {"id": p.get("id") or "", "name": ((p.get("displayName") or {}).get("text") or "").strip(),
                 "domain": domain_of(own_site(p.get("websiteUri") or "")) if p.get("websiteUri") else "",
                 "reviews": int(p.get("userRatingCount") or 0)} for p in places]}
            cache["searches"][query] = hit
        for f in hit["firms"]:
            ident = f["id"] or f["name"].lower()
            if ident not in seen:
                seen.add(ident)
                firms.append(f)
    return firms


def score_town(town: str, firms: list[dict], known: set[str], psi_key: str, cache: dict, score=None) -> dict:
    from google_extras import mobile_score

    score = score or mobile_score
    established = [f for f in firms if f["reviews"] >= ESTABLISHED]
    with_site = [f for f in established if f["domain"]]
    to_measure = sorted(with_site, key=lambda f: -f["reviews"])[:SAMPLE]
    need = [f["domain"] for f in to_measure
            if not _fresh((cache["scores"].get(f["domain"]) or {}).get("at", ""))]
    if need and psi_key:
        with ThreadPoolExecutor(max_workers=4) as pool:
            for d, s in zip(need, pool.map(lambda d: score(d, psi_key), need)):
                cache["scores"][d] = {"at": date.today().isoformat(), "score": s}
    scores = [cache["scores"][f["domain"]]["score"] for f in to_measure
              if isinstance((cache["scores"].get(f["domain"]) or {}).get("score"), int)]
    slow_share = sum(1 for s in scores if s < SLOW) / len(scores) if scores else 0.0
    new_with_site = [f for f in with_site if f["domain"] not in known]
    no_site = [f for f in established if not f["domain"]]
    already = sum(1 for f in firms if f["domain"] and f["domain"] in known)
    return {
        "Town": town,
        "Opportunity": round(len(new_with_site) * slow_share + 0.5 * len(no_site), 1),
        "Firms on Google Maps": len(firms),
        "Established (10+ reviews)": len(established),
        "With a website": len(with_site),
        "Speed checked": len(scores),
        "Median score": round(statistics.median(scores)) if scores else None,
        "Under 50": f"{round(100 * slow_share)}%" if scores else "-",
        "Established, no website": len(no_site),
        "Already in your lists": already,
        "Busiest firm's reviews": max((f["reviews"] for f in firms), default=0),
    }


def rank(rows: list[dict]) -> list[dict]:
    rows = sorted(rows, key=lambda r: (-r["Opportunity"], -r["Established (10+ reviews)"], r["Town"]))
    return [{"Rank": i, **r} for i, r in enumerate(rows, 1)]


def write(outreach: Path, rows: list[dict]) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Areas"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c) for c in COLUMNS])
    ws.freeze_panes = "C2"
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(10, min(26, len(c) + 2))
    notes = wb.create_sheet("How it's worked out")
    for line in __doc__.strip().splitlines():
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    path = outreach / OUT
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)
    return path


def report(rows: list[dict], trades: list[str]) -> list[str]:
    out = [f"Area scout ({', '.join(trades)}): {len(rows)} towns, best first.", "",
           f"{'#':>2}  {'Town':<16} {'Opp.':>5}  {'Busy firms':>10}  {'Median score':>12}  {'Under 50':>8}  {'No site':>7}"]
    for r in rows[:15]:
        median = "-" if r["Median score"] is None else str(r["Median score"])
        out.append(f"{r['Rank']:>2}  {r['Town'][:16]:<16} {r['Opportunity']:>5}  {r['Established (10+ reviews)']:>10}  "
                   f"{median:>12}  {r['Under 50']:>8}  {r['Established, no website']:>7}")
    if rows:
        top = ", ".join(r["Town"] for r in rows[:3])
        out += ["", f"Next: Find new prospects -> Google Maps, the same trades, Areas: {top}."]
    return out


def run(towns: list[str], trades: list[str], outreach: Path, places_key: str, psi_key: str, post=None, score=None,
        log=print) -> list[dict]:
    from find_prospects import known_firms

    if not places_key:
        raise SystemExit("Add GOOGLE_PLACES_API_KEY in Settings (or a PAGESPEED_API_KEY with Places API (New) enabled).")
    if not psi_key:
        log("No PAGESPEED_API_KEY - towns are ranked on firm counts only, without speed scores.")
    cache = load_cache(outreach)
    known = known_firms(outreach)[2]
    rows = []
    try:
        for i, town in enumerate(towns, 1):
            firms = firms_in(town, trades, places_key, cache, post)
            row = score_town(town, firms, known, psi_key, cache, score)
            rows.append(row)
            median = "-" if row["Median score"] is None else row["Median score"]
            log(f"[{i}/{len(towns)}] {town}: {row['Established (10+ reviews)']} established firms, median score {median}")
            save_cache(outreach, cache)  # stopping part-way loses nothing
    finally:
        save_cache(outreach, cache)
    ranked = rank(rows)
    path = write(outreach, ranked)
    for line in report(ranked, trades):
        log(line)
    log(f"Saved {path.name} (every town, every column).")
    return ranked


def main(argv: list[str] | None = None) -> int:
    from find_prospects import TRADES, split_list
    from places_finder import PlacesError

    ap = argparse.ArgumentParser(description="Rank towns by how many busy trade firms have slow websites.")
    ap.add_argument("--trades", default="roofing")
    ap.add_argument("--areas", default="", help="towns, comma separated (blank: South East commuter towns)")
    ap.add_argument("--outreach", type=Path, default=next(
        (d / "outreach" for d in [HERE.parent, *HERE.parents] if (d / "outreach").is_dir()), HERE.parent / "outreach"))
    args = ap.parse_args(argv)
    trades = [t.strip() for t in args.trades.split(",") if t.strip() in TRADES]
    if not trades:
        raise SystemExit("Tick at least one trade.")
    towns = [t.strip().title() for t in split_list(args.areas)] or DEFAULT_TOWNS
    args.outreach.mkdir(parents=True, exist_ok=True)
    places_key = (os.environ.get("GOOGLE_PLACES_API_KEY") or os.environ.get("PAGESPEED_API_KEY") or "").strip()
    try:
        run(towns, trades, args.outreach, places_key, (os.environ.get("PAGESPEED_API_KEY") or "").strip())
    except PlacesError as e:
        raise SystemExit(str(e)) from e
    return 0


if __name__ == "__main__":
    sys.exit(main())
