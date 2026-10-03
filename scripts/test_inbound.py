"""inbound: a website lead into the pipeline - parsed safely, never into a cold batch."""

import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import email_batches
import inbound
import telegram_bot
from test_telegram_bot import CHAT, fake_panel

ALERT = """🔥 Inbound lead - asked for their website report (free speed test)
Reply today: they asked to hear from you.

Business: Kerr Roofing
Name: John Kerr
Email: john@kerrroofing.co.uk
Phone: 07401 111222
Website: https://kerrroofing.co.uk
Score: 34/100 on mobile
Trade: Roofing
Town: Guildford"""


class Parse(unittest.TestCase):
    def test_reads_the_alert(self):
        lead = inbound.parse(ALERT)
        self.assertEqual((lead["business"], lead["email"], lead["website"], lead["score"], lead["town"]),
                         ("Kerr Roofing", "john@kerrroofing.co.uk", "https://kerrroofing.co.uk", "34", "Guildford"))

    def test_not_a_lead(self):
        self.assertIsNone(inbound.parse("Done: Run the whole list"))
        self.assertIsNone(inbound.parse(ALERT.replace("john@kerrroofing.co.uk", "nope")))

    def test_no_formulas(self):
        self.assertEqual(inbound.safe("=HYPERLINK(1)"), "'=HYPERLINK(1)")


class Add(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def test_added_once(self):
        lead = inbound.parse(ALERT)
        self.assertEqual(inbound.add(self.d, lead)[0], "added")
        self.assertEqual(inbound.add(self.d, lead)[0], "already")
        rows = list(openpyxl.load_workbook(self.d / inbound.SHEET)["Outreach"].iter_rows(values_only=True))
        header, row = rows[0], dict(zip(rows[0], rows[1]))
        self.assertEqual(len(rows), 2)
        self.assertIn("Source", header)
        self.assertEqual((row["Email"], row["Contact name"], row["Status"], row["Their score"]), ("john@kerrroofing.co.uk", "John Kerr", "New", "34"))

    def test_known_in_another_list(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Outreach"
        wb.active.append(["Business", "Website"])
        wb.active.append(["Kerr Roofing", "www.kerrroofing.co.uk"])
        wb.save(self.d / "roofers.xlsx")
        status, msg = inbound.add(self.d, inbound.parse(ALERT))
        self.assertEqual(status, "known:roofers.xlsx")
        self.assertFalse((self.d / inbound.SHEET).exists())

    def test_never_in_a_cold_batch(self):
        for name in ("mailmeteor-inbound.csv", "mailmeteor-roofers.csv"):
            (self.d / name).write_text("business,email\nA,a@a.co.uk\n", encoding="utf-8")
        self.assertEqual([p.name for p in email_batches.sources(self.d, None)], ["mailmeteor-roofers.csv"])


class BotButton(unittest.TestCase):
    def test_add_to_pipeline_runs_the_list_and_offers_a_reply(self):
        d = Path(tempfile.mkdtemp())
        panel, _ = fake_panel(d)
        panel.load_settings = lambda: {"PROSPECTS_API_SECRET": "x" * 40, "SITE_URL": "https://www.scalardigital.co.uk"}
        panel.ACTIONS["all"] = "Run the whole list"
        sent = []
        bot = telegram_bot.Bot(panel, "tok", CHAT, http=lambda m, p, timeout=30: sent.append((m, p)) or {})
        with mock.patch("autopilot.is_running", return_value=False), mock.patch.object(telegram_bot.threading, "Thread") as thread:
            bot.handle({"callback_query": {"id": "q", "data": "inbound", "message": {"chat": {"id": int(CHAT)}, "text": ALERT}}})
        self.assertEqual(panel.JOB.started, ["Run the whole list"])
        self.assertTrue((d / inbound.SHEET).exists())
        then = thread.call_args.kwargs["args"][0]
        self.assertIn("/for/kerr-roofing-", then[0])
        reply = next(v for v in bot.items.values() if v["kind"] == "reply")["item"]
        self.assertEqual((reply["email"], reply["sheet"]), ("john@kerrroofing.co.uk", inbound.SHEET))
        self.assertIn("free speed test and scored 34/100", reply["message"])

    def test_a_stranger_cant_add_leads(self):
        d = Path(tempfile.mkdtemp())
        panel, _ = fake_panel(d)
        bot = telegram_bot.Bot(panel, "tok", CHAT, http=lambda m, p, timeout=30: {})
        bot.handle({"callback_query": {"id": "q", "data": "inbound", "message": {"chat": {"id": 999}, "text": ALERT}}})
        self.assertFalse((d / inbound.SHEET).exists())


if __name__ == "__main__":
    unittest.main()
