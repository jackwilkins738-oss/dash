"""data_request: find, export and erase one prospect everywhere on this PC and the dashboard."""

import csv
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import data_request


def write(path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def read(path):
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


class DataRequest(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outreach"
        ws.append(["Business", "Website", "Email", "Status"])
        ws.append(["Kerr Roofing", "https://www.kerrroofing.co.uk/", "info@kerrroofing.co.uk", "Emailed"])
        ws.append(["Other Roofing", "other.co.uk", "hi@other.co.uk", "New"])
        wb.save(self.d / "list.xlsx")
        write(self.d / "emails-sent.csv", ["email", "business", "preview_url"], [
            {"email": "info@kerrroofing.co.uk", "business": "Kerr Roofing", "preview_url": "https://x/for/kerr-roofing-abc123?src=email"},
            {"email": "hi@other.co.uk", "business": "Other Roofing", "preview_url": "https://x/for/other-roofing-def456"}])
        write(self.d / "teardown-log.csv", ["website", "result"], [{"website": "kerrroofing.co.uk", "result": "ok"}, {"website": "other.co.uk", "result": "ok"}])
        write(self.d / "calls.csv", ["key", "outcome", "note"], [{"key": "x", "outcome": "spoke", "note": "see /for/kerr-roofing-abc123"}])

    def test_bad_input(self):
        with self.assertRaises(data_request.RequestError):
            data_request.Who("Kerr Roofing")

    def test_found_by_website_learns_email_and_preview(self):
        who = data_request.Who("kerrroofing.co.uk")
        found = data_request.find(self.d, who)
        self.assertEqual(sorted(found), ["calls.csv", "emails-sent.csv", "list.xlsx (Outreach)", "teardown-log.csv"])
        self.assertIn("info@kerrroofing.co.uk", who.emails)
        self.assertIn("kerr-roofing-abc123", who.slugs)
        self.assertTrue(all(len(rows) == 1 for rows in found.values()))

    def test_export_writes_a_plain_report_with_the_dashboard_record(self):
        with mock.patch.object(data_request, "dashboard", return_value={"prospect": {"business_name": "Kerr Roofing", "view_count": 3, "teardown": {"big": 1}}}):
            path = data_request.export(self.d, "info@kerrroofing.co.uk", "https://dash", "s" * 40)
        text = path.read_text(encoding="utf-8")
        self.assertIn("about Kerr Roofing", text)
        self.assertIn("view_count: 3", text)
        self.assertNotIn("big", text)
        self.assertNotIn("Other Roofing", text)

    def test_erase_removes_them_everywhere_and_keeps_only_the_suppression_line(self):
        calls = []

        def fake(api, secret, slug, method="GET"):
            calls.append((slug, method))
            return {"ok": True, "deleted": True}

        with mock.patch.object(data_request, "dashboard", fake):
            notes = data_request.erase(self.d, "kerrroofing.co.uk", "https://dash", "s" * 40)
        self.assertIn(("kerr-roofing-abc123", "DELETE"), calls)
        self.assertEqual([r["email"] for r in read(self.d / "emails-sent.csv")], ["hi@other.co.uk"])
        self.assertEqual([r["website"] for r in read(self.d / "teardown-log.csv")], ["other.co.uk"])
        self.assertEqual(read(self.d / "calls.csv"), [])
        ws = openpyxl.load_workbook(self.d / "list.xlsx")["Outreach"]
        self.assertEqual([r[0] for r in ws.iter_rows(min_row=2, values_only=True)], ["Other Roofing"])
        block = {(r["kind"], r["value"]) for r in read(self.d / "do-not-contact.csv")}
        self.assertEqual(block, {("website", "kerrroofing.co.uk"), ("email", "info@kerrroofing.co.uk")})
        self.assertTrue(any("preview record" in n for n in notes))

    def test_a_dashboard_refusal_changes_nothing_here(self):
        with mock.patch.object(data_request, "dashboard", side_effect=data_request.RequestError("client")):
            with self.assertRaises(data_request.RequestError):
                data_request.erase(self.d, "kerrroofing.co.uk", "https://dash", "s" * 40)
        self.assertEqual(len(read(self.d / "emails-sent.csv")), 2)
        self.assertFalse((self.d / "do-not-contact.csv").exists())


if __name__ == "__main__":
    unittest.main()
