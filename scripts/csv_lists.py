"""A list saved as CSV (a download from a lead site, an export, Excel's "Save as CSV") works like any other.

The pipeline reads .xlsx workbooks with an "Outreach" tab. Any CSV dropped in outreach/ that looks like a
list of firms - a name column and a website column, under any of the usual headings - gets a matching
<same name>.xlsx made beside it, which then shows in the panel's list picker and runs like the rest.

The CSV itself is never changed. Download a newer copy over it and the firms it adds (by website) are
appended to the workbook; firms already there are left alone, so statuses and notes the pipeline keeps
are never overwritten. The panel's own CSVs (mailmeteor-*, preview-links-* and the rest) never match:
they have no business-name + website pair under these headings, and their names are skipped anyway.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

# Workbook heading -> the CSV headings that mean it (compared lower case, letters and digits only).
ALIASES = {
    "Business": ["business", "businessname", "name", "company", "companyname", "title", "tradingname"],
    "Website": ["website", "site", "url", "websiteurl", "web", "domain", "homepage"],
    "Email": ["email", "emails", "emailaddress", "email1", "contactemail"],
    "Phone": ["phone", "phonenumber", "telephone", "tel", "mobile", "phone1"],
    "Area": ["area", "city", "town", "location", "locality"],
    "Trade": ["trade", "category", "type", "industry", "maincategory"],
    "Contact name": ["contactname", "contact", "owner", "ownername", "fullname"],
    "Address": ["address", "fulladdress", "postaladdress", "streetaddress"],
    "Status": ["status"],
    "Company type": ["companytype"],
}
OURS = re.compile(r"^(mailmeteor|preview-links|emails-sent|email-checks|teardown-log|letters|contacts-found|claims|"
                  r"do-not-contact|first-lines|overrides|guesses|client-|calls|replies|objections|followups?)", re.I)


def _key(heading: str) -> str:
    return re.sub(r"[^a-z0-9]", "", heading.lower())


def columns(header: list[str]) -> dict[str, str] | None:
    """Workbook heading -> CSV heading, or None when this CSV isn't a list of firms."""
    by_key = {_key(h): h for h in header if h and h.strip()}
    found = {}
    for target, names in ALIASES.items():
        for n in names:
            if n in by_key:
                found[target] = by_key[n]
                break
    return found if "Business" in found and "Website" in found else None


def _domain(url: str) -> str:
    d = re.sub(r"^[a-z]+://", "", (url or "").strip().lower()).split("/", 1)[0]
    return d.removeprefix("www.")


def _read(path: Path) -> tuple[list[str], list[dict]]:
    for enc in ("utf-8-sig", "cp1252"):  # Excel saves "CSV" as Windows-1252 unless told UTF-8
        try:
            with path.open(encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                return list(reader.fieldnames or []), list(reader)
        except UnicodeDecodeError:
            continue
    return [], []


def import_one(path: Path) -> int | None:
    """Makes or tops up <stem>.xlsx from one CSV -> firms added, or None when it isn't a list."""
    import openpyxl
    from find_prospects import HEADERS

    header, rows = _read(path)
    cols = columns(header)
    if cols is None:
        return None
    heads = HEADERS + [h for h in cols if h not in HEADERS]
    book = path.with_suffix(".xlsx")
    if book.exists():
        wb = openpyxl.load_workbook(book)
        if "Outreach" not in wb.sheetnames:
            return None  # someone's own workbook with the same name - never touched
        ws = wb["Outreach"]
        existing = [str(c.value or "").strip() for c in ws[1]]
        site_col = existing.index("Website") if "Website" in existing else None
        have = {_domain(str(r[site_col] or "")) for r in ws.iter_rows(min_row=2, values_only=True)} if site_col is not None else set()
        heads = existing
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outreach"
        ws.append(heads)
        have = set()
    added = 0
    for r in rows:
        name, site = str(r.get(cols["Business"]) or "").strip(), str(r.get(cols["Website"]) or "").strip()
        d = _domain(site)
        if not name or not d or d in have:
            continue
        have.add(d)
        values = {target: str(r.get(src) or "").strip() for target, src in cols.items()}
        values.setdefault("Status", "New")
        values["Status"] = values["Status"] or "New"
        ws.append([values.get(h, "") for h in heads])
        added += 1
    if added or not book.exists():
        tmp = book.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        tmp.replace(book)
    return added


_SEEN: dict[str, float] = {}


def import_all(outreach: Path) -> list[str]:
    """Every list-shaped CSV in outreach/ made into (or topped up in) its workbook -> lines for the log.
    Each CSV is only read again after it changes, so this is cheap to call on every page load."""
    notes = []
    if not outreach.is_dir():
        return notes
    for path in sorted(outreach.glob("*.csv")):
        if OURS.match(path.name):
            continue
        try:
            stamp = path.stat().st_mtime
            if _SEEN.get(str(path)) == stamp:
                continue
            added = import_one(path)
            _SEEN[str(path)] = stamp
        except (OSError, ValueError, KeyError) as e:  # open in Excel, a corrupt file: try again next time
            notes.append(f"Couldn't read {path.name} as a list ({type(e).__name__}) - close it in Excel and refresh.")
            continue
        if added:
            notes.append(f"{path.name}: {added} firms added to {path.with_suffix('.xlsx').name}.")
    return notes
