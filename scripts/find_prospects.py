"""Build a new prospect list from Companies House: trade + area in, a ready sheet out.

Used from the control panel's "Find new prospects" card, or directly:

    python scripts/find_prospects.py --trades roofing --areas "Guildford, Woking" --max 100
    python scripts/find_prospects.py --trades roofing,lofts --areas GU1 --count-only

For every active limited company registered in the areas under the trades'
industry codes, it:
  1. skips firms already in any sheet in outreach/ (so repeat runs give the
     next batch, and nobody already contacted - or who said no - comes back);
  2. looks for their website by trying the domains their name suggests
     (smithroofing.co.uk, .com ...), and keeps one only if the site proves it's
     theirs: their company number, or their name plus something local (their
     registered postcode district or town). A site with just the name - maybe
     a namesake elsewhere - goes on a "Check website" tab for you to confirm;
  3. reads the published business email and phone number from that site;
  4. takes the contact name from the register's list of directors.

It writes a NEW workbook in outreach/, never touching existing sheets, with
three tabs: Outreach (website confirmed - ready for the panel's other
buttons), Check website, and No website (letter prospects - often the best
ones for a web designer). Every firm is a limited company, so the Company
type column is filled and email is allowed (PECR corporate subscribers).

Env:
    COMPANIES_HOUSE_API_KEY  required (free REST key)
"""

from __future__ import annotations

import argparse
import html as htmllib
import re
import socket
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlencode, urlparse

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from company_lookup import (  # noqa: E402 - after the path setup
    CORPORATE_TYPES,
    GENERIC_WORDS,
    LookupFailed,
    _get,
    district,
    key_problem,
    names_called_ltd,
    normalise,
    postcodes,
    site_text,
    useful_links,
)

# ---------------------------------------------------------------- trades

# Industry (SIC) codes on the register, and - where a code is broader than the
# trade - words the company name must contain. Most specific first: a firm that
# matches two trades gets the first.
TRADES: dict[str, dict] = {
    "lofts": {
        "label": "Loft conversions",
        "sic": ["41202", "43390", "43999", "43910", "43320"],
        "name_any": ["loft", "dormer", "attic"],
    },
    "driveways": {
        "label": "Driveways & patios",
        "sic": ["43999", "42110", "43120", "81300", "42990"],
        "name_any": ["driveway", "paving", "resin", "surfacing", "patio", "tarmac", "block"],
    },
    "roofing": {"label": "Roofing", "sic": ["43910"], "name_any": []},
    "landscaping": {"label": "Landscaping", "sic": ["81300"], "name_any": []},
    "building": {"label": "Building & extensions", "sic": ["41202", "43390", "41201"], "name_any": []},
}

AGE_BANDS = {
    "any": (None, None),
    # Registered in the last 6 months: the firms most likely to have no website
    # yet (they land on the No website tab, for a letter), and to want one.
    "under6m": (0.5, None),
    "under2": (2, None),  # incorporated in the last 2 years
    "2to10": (10, 2),
    "over10": (None, 10),
}

DEFAULT_EXCLUDE = "holdings, investments, estates, lettings, capital, finance"

# ---------------------------------------------------------------- register search


def years_ago(years: float, today: date | None = None) -> str:
    today = today or date.today()
    if years != int(years):  # part-years (6 months) - by days, near enough for a search filter
        from datetime import timedelta

        return (today - timedelta(days=round(years * 365.25))).isoformat()
    years = int(years)
    try:
        return today.replace(year=today.year - years).isoformat()
    except ValueError:  # 29 February
        return today.replace(year=today.year - years, day=28).isoformat()


def search_companies(api_key: str, sic: list[str], area: str, age: str, log) -> list[dict]:
    """Active limited companies registered in this area under these industry codes.

    One request per code (and page), so nothing depends on how the API reads a list."""
    newer_than, older_than = AGE_BANDS[age]
    base = {"company_status": "active", "location": area, "size": 500}
    if newer_than:
        base["incorporated_from"] = years_ago(newer_than)
    if older_than:
        base["incorporated_to"] = years_ago(older_than)
    items: dict[str, dict] = {}
    for code in sic:
        start = 0
        while start < 5000:
            page = _get("/advanced-search/companies?" + urlencode({**base, "sic_codes": code, "start_index": start}), api_key)
            batch = page.get("items") or []
            for item in batch:
                # Limited companies and LLPs only: the register also lists partnerships and others.
                if item.get("company_type") in CORPORATE_TYPES and item.get("company_number"):
                    items.setdefault(item["company_number"], item)
            hits = page.get("hits")
            if not batch or len(batch) < base["size"] or (hits is not None and start + len(batch) >= hits):
                break
            start += len(batch)
    log(f"  {area}: {len(items)} companies on the register")
    return list(items.values())


def words_in(text: str, words: list[str]) -> bool:
    lowered = f" {normalise(text)} "
    return any(w.strip() and w.strip().lower() in lowered for w in words)


def trade_for(item: dict, wanted: list[str]) -> str | None:
    """Which of the wanted trades this company is, or None."""
    codes = set(item.get("sic_codes") or [])
    name = item.get("company_name") or ""
    for key in TRADES:  # most specific first
        if key not in wanted:
            continue
        t = TRADES[key]
        if codes & set(t["sic"]) and (not t["name_any"] or words_in(name, t["name_any"])):
            return key
    return None


def split_list(text: str) -> list[str]:
    return [w.strip() for w in re.split(r"[,;\n]", text or "") if w.strip()]


# ---------------------------------------------------------------- already known


def known_firms(outreach: Path) -> tuple[set[str], set[str], set[str]]:
    """Company numbers, names and website domains in every sheet (every tab) in outreach/."""
    from push_prospects import domain_of

    import openpyxl

    numbers, names, domains = set(), set(), set()
    for path in outreach.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = [str(h).strip() if h else "" for h in next(rows, ())]
            for values in rows:
                row = dict(zip(header, values))
                if row.get("Company number"):
                    numbers.add(str(row["Company number"]).strip().upper().zfill(8))
                if row.get("Business"):
                    names.add(normalise(str(row["Business"])))
                for col in ("Website", "Possible website"):
                    d = domain_of(str(row.get(col) or ""))
                    if d:
                        domains.add(d)
        wb.close()
    # Firms in older sheets have no number column, but the company lookup may have found theirs.
    lookups = outreach / "company-lookups.csv"
    if lookups.exists():
        import csv

        with lookups.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r.get("number"):
                    numbers.add(r["number"].strip().upper().zfill(8))
                if r.get("website"):
                    domains.add(r["website"].strip().lower())
    names.discard("")
    return numbers, names, domains


# ---------------------------------------------------------------- website


# Where a guessed domain sometimes lands that isn't the firm's own site.
NOT_THEIR_SITE = re.compile(
    r"checkatrade|facebook|yell\.|mybuilder|trustatrader|ratedpeople|bark\.com|google\.|instagram|linkedin|"
    r"houzz|nextdoor|thomsonlocal|freeindex|cylex|192\.com|companieshouse|endole|opencorporates",
    re.I,
)
PARKED = re.compile(
    r"domain (?:is|may be) for sale|buy this domain|this domain is parked|parked free|domain parking|"
    r"hugedomains|sedo\.com|dan\.com|future home of|coming soon|under construction|account suspended",
    re.I,
)
TRAILING_GENERIC = {"services", "solutions", "contractors", "group", "uk", "and", "sons", "son", "southern", "south"}


def domain_candidates(company_name: str) -> list[str]:
    """Domains the name suggests, likeliest first: "J Smith Roofing Services Ltd" ->
    smithroofing.co.uk, jsmithroofing.co.uk, smithroofingservices.co.uk ..."""
    words = normalise(company_name).split()
    if not words:
        return []

    def no_initials(ws: list[str]) -> list[str]:
        return [w for i, w in enumerate(ws) if not (len(w) == 1 and i < 2)] or ws

    trimmed = list(words)
    while len(trimmed) > 1 and trimmed[-1] in TRAILING_GENERIC:
        trimmed.pop()
    forms = [
        no_initials(trimmed),  # smith roofing
        trimmed,  # j smith roofing
        no_initials(words),  # smith roofing services
        words,  # j smith roofing services
        [w for w in words if w != "and"],  # smith sons (from "smith and sons")
    ]
    joined: list[str] = []
    for f in forms:
        for name in ("".join(f), "-".join(f)):
            if 3 <= len(name) <= 50 and name not in joined:
                joined.append(name)
    # Unhyphenated first: far more common for small firms.
    joined.sort(key=lambda n: "-" in n)
    names = joined[:8]
    return [n + ".co.uk" for n in names] + [n + ".com" for n in names] + [n + ".uk" for n in names[:4]]


def resolves(domain: str) -> bool:
    try:
        socket.getaddrinfo(domain, 443, proto=socket.IPPROTO_TCP)
        return True
    except (socket.gaierror, UnicodeError, OSError):
        return False


def site_domain(url: str | None) -> str:
    host = (urlparse(url or "").hostname or "").lower()
    return host[4:] if host.startswith("www.") else host


def judge_site(item: dict, pages_html: str, area: str) -> tuple[str, str]:
    """(confirmed | possible | no, why) - is this site this company's?"""
    text = site_text(pages_html)
    if len(text) < 200 or PARKED.search(text[:3000]):
        return "no", "parked or empty"
    number = (item.get("company_number") or "").upper()
    digits = number.lstrip("0") if number.isdigit() else number
    if number and re.search(rf"\b0*{re.escape(digits)}\b", text):
        return "confirmed", "company number on their site"
    name = normalise(item.get("company_name") or "")
    flat = f" {normalise(text)} "
    named = bool(name) and f" {name} " in flat
    if not named:
        # "J Smith Roofing Ltd" often calls itself "Smith Roofing": the distinctive words, together.
        distinctive = [w for w in name.split() if len(w) > 1 and w not in GENERIC_WORDS]
        core = " ".join(w for w in name.split() if len(w) > 1)
        named = bool(distinctive) and f" {core} " in flat
    if not named:
        return "no", "doesn't name them"
    office = item.get("registered_office_address") or {}
    reg_pc = (office.get("postal_code") or "").replace(" ", "").upper()
    if reg_pc and district(reg_pc) in {district(p) for p in postcodes(text)}:
        return "confirmed", "their name + a postcode near their registered office"
    for town in {office.get("locality") or "", area}:
        town = town.strip()
        if len(town) >= 4 and not re.fullmatch(r"[A-Z]{1,2}\d.*", town.upper()) and re.search(rf"\b{re.escape(town)}\b", text, re.I):
            return "confirmed", f"their name + {town.title()} on their site"
    if normalise(item.get("company_name") or "") in names_called_ltd(text):
        return "confirmed", "their site calls them by their registered name"
    return "possible", "their name is on it, but nothing local - could be a namesake"


def find_website(item: dict, area: str, fetch, resolve=None) -> dict:
    """The company's own website, if its name leads to one: {status, website, how, pages}."""
    resolve = resolve or resolves
    possible = None
    for domain in domain_candidates(item.get("company_name") or ""):
        if not resolve(domain):
            continue
        html, final = fetch(f"https://{domain}/")
        if html is None:
            html, final = fetch(f"http://{domain}/")
        if html is None:
            continue
        landed = site_domain(final) or domain
        if NOT_THEIR_SITE.search(landed):
            continue
        pages = [html]
        for url in useful_links(html, final or f"https://{domain}/", limit=3):
            extra, _ = fetch(url)
            if extra:
                pages.append(extra)
        joined = "\n".join(pages)
        status, how = judge_site(item, joined, area)
        if status == "confirmed":
            return {"status": "confirmed", "website": landed, "how": how, "pages": joined}
        if status == "possible" and possible is None:
            possible = {"status": "possible", "website": landed, "how": how, "pages": joined}
    return possible or {"status": "none", "website": "", "how": "no site found from their name", "pages": ""}


# ---------------------------------------------------------------- contact details

EMAIL = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
JUNK_EMAIL = re.compile(
    r"example\.|sentry|wixpress|domain\.com|email\.com|yourname|name@|user@|@2x|\.(png|jpe?g|gif|webp|svg)$|"
    r"godaddy|squarespace|wordpress|jquery|bootstrap|u003e|noreply|no-reply",
    re.I,
)
PREFERRED_BOX = ["info", "enquiries", "enquiry", "office", "sales", "contact", "hello", "admin", "quotes", "mail"]
FREE_MAIL = {"gmail.com", "googlemail.com", "hotmail.com", "hotmail.co.uk", "outlook.com", "live.co.uk", "yahoo.co.uk",
             "yahoo.com", "btinternet.com", "icloud.com", "aol.com", "sky.com", "live.com", "msn.com"}


def decode_cloudflare(hexstr: str) -> str:
    """Cloudflare's "email protection" hides addresses as XOR'd hex; this reverses it."""
    try:
        key = int(hexstr[:2], 16)
        return "".join(chr(int(hexstr[i : i + 2], 16) ^ key) for i in range(2, len(hexstr), 2))
    except ValueError:
        return ""


def find_email(pages_html: str, site: str) -> str:
    """The business's published email: one on its own domain first, a preferred inbox first."""
    raw = htmllib.unescape(pages_html)
    found = [m.group(0) for m in EMAIL.finditer(raw)]
    found += [decode_cloudflare(h) for h in re.findall(r"data-cfemail=[\"']([0-9a-fA-F]+)", raw)]
    found += [decode_cloudflare(h) for h in re.findall(r"email-protection#([0-9a-fA-F]+)", raw)]
    seen: list[str] = []
    for e in found:
        e = e.strip().strip(".").lower()
        if EMAIL.fullmatch(e) and not JUNK_EMAIL.search(e) and e not in seen:
            seen.append(e)
    if not seen:
        return ""

    def rank(e: str) -> tuple:
        box, dom = e.split("@", 1)
        own = site and (dom == site or dom.endswith("." + site))
        return (0 if own else 1 if dom in FREE_MAIL else 2, PREFERRED_BOX.index(box) if box in PREFERRED_BOX else 99)

    best = sorted(seen, key=rank)[0]
    return best if rank(best)[0] < 2 else ""  # someone else's domain: a supplier, not them


PHONE = re.compile(r"(?:\+44\s?\(?0?\)?|\b0)(?:\d[\s-]?){9,10}\b")


def find_phone(pages_html: str) -> str:
    raw = htmllib.unescape(pages_html)
    tel = re.findall(r"href=[\"']tel:([^\"']+)", raw, re.I)
    candidates = tel + [m.group(0) for m in PHONE.finditer(site_text(raw))]
    for c in candidates:
        digits = re.sub(r"\D", "", c)
        if digits.startswith("44"):
            digits = "0" + digits[2:].lstrip("0")
        if len(digits) == 11 and digits.startswith(("01", "02", "03", "07", "08")):
            return f"{digits[:5]} {digits[5:]}" if digits.startswith(("01", "07", "08")) else f"{digits[:3]} {digits[3:7]} {digits[7:]}"
    return ""


def pick_director(officers: dict) -> tuple[str, int]:
    """(contact name, how many active directors) - the longest-serving director, as "John Smith"."""
    people = [
        o for o in officers.get("items") or []
        if not o.get("resigned_on") and o.get("officer_role") in ("director", "llp-designated-member", "llp-member")
    ]
    if not people:
        return "", 0
    people.sort(key=lambda o: o.get("appointed_on") or "9999")
    raw = people[0].get("name") or ""
    if "," in raw:
        surname, forenames = raw.split(",", 1)
        first = forenames.split()[0] if forenames.split() else ""
        name = f"{first} {surname}".strip()
    else:
        name = raw
    name = " ".join(name.lower().split()).title()  # "o'neill-smith" -> "O'Neill-Smith"
    name = re.sub(r"\bMc([a-z])", lambda m: "Mc" + m.group(1).upper(), name)
    return name, len(people)


# ---------------------------------------------------------------- workbook

HEADERS = [
    "Business", "Website", "Status", "Trade", "Area", "Email", "Phone", "Company type", "Contact name",
    "Company number", "Registered name", "Registered address", "Incorporated", "Directors", "Website found",
    "Mobile score", "LCP (s)",
]
TABS = {"confirmed": "Outreach", "possible": "Check website", "none": "No website"}


def display_name(registered: str) -> str:
    """"J SMITH ROOFING (UK) LIMITED" -> "J Smith Roofing (UK)": how the preview page and emails will name them."""
    name = re.sub(r"\s+", " ", registered).strip()
    name = re.sub(r"[\s,.]+(?:LTD|LIMITED|LLP|PLC|L\.L\.P)\.?$", "", name, flags=re.I)
    name = name.title()
    name = re.sub(r"\b(Uk|Gb|Llp)\b", lambda m: m.group(1).upper(), name)
    name = re.sub(r"'S\b", "'s", name)
    return name


def row_for(item: dict, trade: str, site: dict, email: str, phone: str, contact: str, directors: int) -> dict:
    office = item.get("registered_office_address") or {}
    address = ", ".join(
        str(office.get(k)).strip()
        for k in ("address_line_1", "address_line_2", "locality", "region", "postal_code")
        if office.get(k)
    )
    ctype = CORPORATE_TYPES.get(item.get("company_type") or "", "Ltd")
    name = item.get("company_name") or ""
    return {
        "Business": display_name(name),
        "Registered name": name,
        "Website": site["website"] if site["status"] == "confirmed" else "",
        "Possible website": site["website"] if site["status"] == "possible" else "",
        "Status": "New",
        "Trade": TRADES[trade]["label"],
        "Area": (office.get("locality") or "").strip().title() or item.get("_area", ""),
        "Email": email,
        "Phone": phone,
        "Company type": ctype,
        "Contact name": contact,
        "Company number": item.get("company_number") or "",
        "Registered address": address,
        "Incorporated": item.get("date_of_creation") or "",
        "Directors": directors or "",
        "Website found": site["how"],
    }


def write_workbook(path: Path, rows: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for status, tab in TABS.items():
        ws = wb.create_sheet(tab)
        headers = HEADERS if status != "possible" else HEADERS[:2] + ["Possible website"] + HEADERS[2:]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "B2"
        for r in rows:
            if r["_status"] == status:
                ws.append([r.get(h, "") for h in headers])
        widths = {"Business": 34, "Website": 26, "Possible website": 26, "Email": 30, "Registered address": 44, "Website found": 40}
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(h, 14)
    wb.save(path)


def output_path(outreach: Path, trades: list[str], areas: list[str]) -> Path:
    slug = "-".join(trades[:2] + areas[:2]).lower()
    slug = re.sub(r"[^a-z0-9]+", "-", slug).strip("-")[:60] or "prospects"
    base = f"{slug}-{date.today().isoformat()}"
    path = outreach / f"{base}.xlsx"
    n = 2
    while path.exists():
        path = outreach / f"{base}-{n}.xlsx"
        n += 1
    return path


# ---------------------------------------------------------------- main


def outreach_dir() -> Path:
    for parent in [HERE.parent, *HERE.parents]:
        if (parent / "outreach").is_dir():
            return parent / "outreach"
    return HERE.parent / "outreach"


def main(argv: list[str] | None = None) -> None:
    import os

    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--trades", required=True, help="comma list of: " + ", ".join(TRADES))
    ap.add_argument("--areas", required=True, help="towns or postcode districts, comma separated")
    ap.add_argument("--age", default="any", choices=list(AGE_BANDS))
    ap.add_argument("--max", type=int, default=100, help="most firms to add (after skipping known ones)")
    ap.add_argument("--include", default="", help="only names containing one of these words")
    ap.add_argument("--exclude", default=DEFAULT_EXCLUDE, help="skip names containing any of these words")
    ap.add_argument("--no-websites", action="store_true", help="skip the website search (much faster)")
    ap.add_argument("--no-directors", action="store_true", help="skip director names")
    ap.add_argument("--website-only", action="store_true", help="leave out firms with no website found")
    ap.add_argument("--email-only", action="store_true", help="leave out firms with no email found")
    ap.add_argument("--count-only", action="store_true", help="just say how many match")
    args = ap.parse_args(argv)

    trades = [t for t in split_list(args.trades.lower()) if t in TRADES]
    areas = split_list(args.areas)
    if not trades:
        sys.exit("Pick at least one trade: " + ", ".join(TRADES))
    if not areas:
        sys.exit("Give at least one area - a town like Guildford or a postcode district like GU1.")
    if not 1 <= args.max <= 1000:
        sys.exit("--max must be 1-1000.")
    if (args.email_only or args.website_only) and args.no_websites:
        sys.exit("Only firms with an email / a website needs the website search - untick Skip the website search.")
    key = os.environ.get("COMPANIES_HOUSE_API_KEY", "")
    if not key or key_problem(key):
        sys.exit("Needs a COMPANIES_HOUSE_API_KEY (Settings) - " + (key_problem(key) or "not set") + ".")

    outreach = outreach_dir()
    outreach.mkdir(exist_ok=True)
    log = lambda line: print(line, flush=True)  # noqa: E731

    sic = sorted({c for t in trades for c in TRADES[t]["sic"]})
    log(f"Searching Companies House: {', '.join(TRADES[t]['label'] for t in trades)} in {', '.join(areas)} ...")
    found: dict[str, dict] = {}
    try:
        for area in areas:
            for item in search_companies(key, sic, area, args.age, log):
                num = (item.get("company_number") or "").upper()
                if num and num not in found:
                    item["_area"] = area
                    found[num] = item
    except PermissionError as e:
        sys.exit(f"{e}.")
    except LookupFailed as e:
        sys.exit(f"Companies House didn't answer - {e}. Nothing was written.")

    include, exclude = split_list(args.include), split_list(args.exclude)
    numbers, names, domains = known_firms(outreach)
    counts = {"trade": 0, "include": 0, "exclude": 0, "known": 0}
    keep: list[tuple[dict, str]] = []
    for num, item in sorted(found.items(), key=lambda kv: kv[1].get("date_of_creation") or ""):
        trade = trade_for(item, trades)
        name = item.get("company_name") or ""
        if not trade:
            counts["trade"] += 1
        elif include and not words_in(name, include):
            counts["include"] += 1
        elif exclude and words_in(name, exclude):
            counts["exclude"] += 1
        elif num.zfill(8) in numbers or normalise(name) in names:
            counts["known"] += 1
        else:
            keep.append((item, trade))
    log(
        f"{len(found)} companies found; {len(keep)} new to you "
        f"({counts['known']} already in your sheets, {counts['exclude']} excluded by name, "
        f"{counts['trade'] + counts['include']} not the trade you picked)."
    )
    if args.count_only:
        by_trade: dict[str, int] = {}
        for _, t in keep:
            by_trade[TRADES[t]["label"]] = by_trade.get(TRADES[t]["label"], 0) + 1
        for label, n in sorted(by_trade.items(), key=lambda kv: -kv[1]):
            log(f"  {label}: {n}")
        return
    if not keep:
        sys.exit("Nothing new to add - try more areas, another trade, or a wider age range.")
    filtered = args.email_only or args.website_only
    if filtered:
        # "Max" means firms kept: keep checking until there are that many, or no more to check.
        wants = " and ".join(w for w, on in (("a website", args.website_only), ("an email", args.email_only)) if on)
        log(f"Checking up to {len(keep)} firms for {args.max} with {wants} ...")
    else:
        keep = keep[: args.max]
        log(f"Building a list of {len(keep)} ...")

    from site_teardown import fetch_html

    fetch = lambda url: fetch_html(url, timeout=12)  # noqa: E731
    rows: list[dict] = []
    out = output_path(outreach, trades, areas)
    lock = threading.Lock()
    checked = [0]
    left_out = {"known": 0, "filtered": 0}

    def wanted(site: dict, email: str) -> bool:
        if args.website_only and site["status"] != "confirmed":
            return False
        if args.email_only and not email:
            return False
        return True

    def work(item: dict, trade: str) -> dict | str:
        try:
            return build_row(item, trade)
        except Exception as e:  # one odd website never costs the whole list
            site = {"status": "none", "website": "", "how": f"website search failed ({type(e).__name__})", "pages": ""}
            if not wanted(site, ""):
                return "filtered"
            row = row_for(item, trade, site, "", "", "", 0)
            row["_status"] = "none"
            return row

    def build_row(item: dict, trade: str) -> dict | str:
        site = {"status": "none", "website": "", "how": "website search skipped", "pages": ""}
        if not args.no_websites:
            site = find_website(item, item["_area"], fetch)
        # Their site is already in a sheet under another name: already known, leave them out.
        if site["website"] and site["website"] in domains:
            return "known"
        email = find_email(site["pages"], site["website"]) if site["pages"] else ""
        if not wanted(site, email):
            return "filtered"  # decided before asking for directors - no request wasted
        phone = find_phone(site["pages"]) if site["pages"] else ""
        contact, directors = "", 0
        if not args.no_directors:
            try:
                contact, directors = pick_director(_get(f"/company/{item['company_number']}/officers?items_per_page=50", key))
            except (LookupFailed, PermissionError):
                pass
        row = row_for(item, trade, site, email, phone, contact, directors)
        row["_status"] = site["status"]
        return row

    def report(result: dict | str, business: str) -> None:
        with lock:
            checked[0] += 1
            prefix = f"[{len(rows) + (1 if isinstance(result, dict) else 0)}/{args.max if filtered else len(keep)}]"
            if result == "known":
                left_out["known"] += 1
                log(f"{prefix} {business}: skipped - their website is already in your sheets")
            elif result == "filtered":
                left_out["filtered"] += 1
                if not filtered or checked[0] % 10 == 0:
                    log(f"    ... {checked[0]} checked, {len(rows)} kept so far")
            else:
                rows.append(result)
                site = result["Website"] or (f"maybe {result['Possible website']}" if result["Possible website"] else "no website")
                bits = [site] + [b for b in (result["Email"], result["Phone"], result["Contact name"]) if b]
                log(f"{prefix} {result['Business']}: " + " · ".join(bits))
                if len(rows) % 20 == 0:
                    write_workbook(out, rows)

    try:
        with ThreadPoolExecutor(max_workers=6) as pool:
            pending = {}
            queue = iter(keep)
            while True:
                # Keep six firms in flight; stop starting new ones once there are enough.
                while len(pending) < 6 and len(rows) < args.max:
                    nxt = next(queue, None)
                    if nxt is None:
                        break
                    pending[pool.submit(work, *nxt)] = nxt[0].get("company_name") or ""
                if not pending:
                    break
                done_now = next(as_completed(pending))
                business = display_name(pending.pop(done_now))
                if len(rows) >= args.max:
                    continue  # enough already: this one's work is dropped
                report(done_now.result(), business)
    except KeyboardInterrupt:
        log("Stopped - saving what's done.")

    rows.sort(key=lambda r: r["Business"])
    write_workbook(out, rows)
    tally = {s: sum(1 for r in rows if r["_status"] == s) for s in TABS}
    with_email = sum(1 for r in rows if r["Email"])
    log("")
    log(f"Wrote {out.name}:")
    log(f"  Outreach       {tally['confirmed']:>4}  website confirmed ({with_email} with an email found overall)")
    log(f"  Check website  {tally['possible']:>4}  a site with their name but nothing local - confirm, then move to Outreach")
    log(f"  No website     {tally['none']:>4}  letter prospects - their registered address is in the sheet")
    if filtered:
        log(f"  ({checked[0]} firms checked; {left_out['filtered']} left out for having no {wants.replace('a ', '').replace('an ', '')})")
        if len(rows) < args.max:
            log(f"  Only {len(rows)} of the {args.max} you asked for - that's every match. Try more areas or trades for more.")
    log("Next: pick it in the List box and press Run the whole list.")


if __name__ == "__main__":
    main()
