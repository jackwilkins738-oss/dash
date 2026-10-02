"""csv_lists: a list saved as CSV shows in the panel and runs like a workbook."""

import csv
import tempfile
import unittest
from pathlib import Path

import openpyxl

import csv_lists


def write(path, fields, rows, encoding="utf-8"):
    with path.open("w", newline="", encoding=encoding) as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def outreach_rows(book):
    ws = openpyxl.load_workbook(book)["Outreach"]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


class CsvLists(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        csv_lists._SEEN.clear()

    def test_a_downloaded_list_becomes_a_workbook(self):
        write(self.d / "lofts-Edinburgh 2026.csv", ["Name", "Site", "Email", "Phone", "City", "Category"], [
            {"Name": "Kerr Lofts", "Site": "https://www.kerrlofts.co.uk/", "Email": "info@kerrlofts.co.uk", "City": "Edinburgh", "Category": "Loft conversion"},
            {"Name": "No Site Lofts", "Site": "", "Email": "a@b.com"},
            {"Name": "Kerr Lofts again", "Site": "kerrlofts.co.uk"},
        ])
        notes = csv_lists.import_all(self.d)
        self.assertIn("lofts-Edinburgh 2026.csv: 1 firms added", notes[0])
        rows = outreach_rows(self.d / "lofts-Edinburgh 2026.xlsx")
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["Business"], rows[0]["Website"], rows[0]["Email"], rows[0]["Area"], rows[0]["Trade"], rows[0]["Status"]),
                         ("Kerr Lofts", "https://www.kerrlofts.co.uk/", "info@kerrlofts.co.uk", "Edinburgh", "Loft conversion", "New"))

    def test_a_newer_download_only_adds_new_firms(self):
        path = self.d / "list.csv"
        write(path, ["Business", "Website"], [{"Business": "A", "Website": "a.co.uk"}])
        csv_lists.import_all(self.d)
        wb = openpyxl.load_workbook(self.d / "list.xlsx")
        wb["Outreach"]["C2"] = "Emailed"  # the pipeline's status - must survive
        wb.save(self.d / "list.xlsx")
        write(path, ["Business", "Website"], [{"Business": "A", "Website": "www.a.co.uk"}, {"Business": "B", "Website": "b.co.uk"}])
        csv_lists._SEEN.clear()
        csv_lists.import_all(self.d)
        rows = outreach_rows(self.d / "list.xlsx")
        self.assertEqual([(r["Business"], r["Status"]) for r in rows], [("A", "Emailed"), ("B", "New")])

    def test_the_panels_own_files_and_other_csvs_are_left_alone(self):
        write(self.d / "mailmeteor-lofts.csv", ["business", "email", "preview_url"], [{"business": "A", "email": "a@a.co.uk"}])
        write(self.d / "notes.csv", ["what", "when"], [{"what": "x", "when": "y"}])
        self.assertEqual(csv_lists.import_all(self.d), [])
        self.assertEqual(sorted(p.name for p in self.d.glob("*.xlsx")), [])

    def test_excel_windows_encoding(self):
        write(self.d / "w.csv", ["Business", "Website"], [{"Business": "Café Roofing", "Website": "cafe.co.uk"}], encoding="cp1252")
        csv_lists.import_all(self.d)
        self.assertEqual(outreach_rows(self.d / "w.xlsx")[0]["Business"], "Café Roofing")

    def test_someone_elses_workbook_with_the_same_name_is_never_touched(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(self.d / "mine.xlsx")
        write(self.d / "mine.csv", ["Business", "Website"], [{"Business": "A", "Website": "a.co.uk"}])
        csv_lists.import_all(self.d)
        self.assertEqual(openpyxl.load_workbook(self.d / "mine.xlsx").sheetnames, ["Sheet1"])


if __name__ == "__main__":
    unittest.main()
