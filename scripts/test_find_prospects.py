"""Offline tests for find_prospects - no network. Run: python -m unittest scripts/test_find_prospects.py"""

import io
import os
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

import find_prospects as fp


def company(name, number="01234567", sic=("43910",), postcode="GU1 4RR", locality="Guildford", created="2015-03-01"):
    return {
        "company_name": name,
        "company_number": number,
        "company_status": "active",
        "company_type": "ltd",
        "date_of_creation": created,
        "sic_codes": list(sic),
        "registered_office_address": {"address_line_1": "1 High St", "locality": locality, "postal_code": postcode},
    }


FILLER = "<p>" + "We carry out roof repairs, re-roofs, flat roofs and chimney work across the county. " * 5 + "</p>"


class Trades(unittest.TestCase):
    def test_specific_trades_win_and_need_their_words(self):
        loft = company("SURREY LOFTS LTD", sic=("41202",))
        builder = company("ACME BUILDERS LTD", sic=("41202",))
        self.assertEqual(fp.trade_for(loft, ["lofts", "building"]), "lofts")
        self.assertEqual(fp.trade_for(builder, ["lofts", "building"]), "building")
        self.assertIsNone(fp.trade_for(builder, ["lofts"]))  # a builder isn't a loft firm without the word

    def test_age_bands(self):
        from datetime import date

        self.assertEqual(fp.years_ago(2, date(2026, 9, 27)), "2024-09-27")
        self.assertEqual(fp.years_ago(1, date(2024, 2, 29)), "2023-02-28")


class Names(unittest.TestCase):
    def test_display_name(self):
        self.assertEqual(fp.display_name("J SMITH ROOFING (UK) LIMITED"), "J Smith Roofing (UK)")
        self.assertEqual(fp.display_name("O'NEILL'S BUILDING SERVICES LTD."), "O'Neill's Building Services")

    def test_director(self):
        officers = {
            "items": [
                {"name": "SMITH, Jane", "officer_role": "director", "appointed_on": "2019-01-01"},
                {"name": "MCDONALD-JONES, John Michael", "officer_role": "director", "appointed_on": "2015-06-01"},
                {"name": "OLD, Person", "officer_role": "director", "appointed_on": "2010-01-01", "resigned_on": "2014-01-01"},
                {"name": "ACME SECRETARIES LTD", "officer_role": "corporate-secretary", "appointed_on": "2009-01-01"},
            ]
        }
        self.assertEqual(fp.pick_director(officers), ("John McDonald-Jones", 2))
        self.assertEqual(fp.pick_director({"items": []}), ("", 0))


class Domains(unittest.TestCase):
    def test_candidates(self):
        c = fp.domain_candidates("J SMITH ROOFING SERVICES LIMITED")
        self.assertEqual(c[0], "smithroofing.co.uk")
        for want in ("smithroofingservices.co.uk", "jsmithroofing.co.uk", "smithroofing.co.uk", "smithroofing.com", "smith-roofing.co.uk"):
            self.assertIn(want, c)
        self.assertLessEqual(len(c), 20)
        self.assertLess(c.index("smithroofing.co.uk"), c.index("smith-roofing.co.uk"))


class JudgeSite(unittest.TestCase):
    def test_number_confirms(self):
        page = f"<footer>Smith Roofing Ltd. Company No. 1234567</footer>{FILLER}"
        self.assertEqual(fp.judge_site(company("SMITH ROOFING LTD"), page, "Guildford")[0], "confirmed")

    def test_name_plus_local_confirms(self):
        item = company("SMITH ROOFING LTD")
        self.assertEqual(fp.judge_site(item, f"<h1>Smith Roofing</h1><p>Guildford GU1 2AB</p>{FILLER}", "")[0], "confirmed")
        self.assertEqual(fp.judge_site(item, f"<h1>Smith Roofing</h1><p>Serving Guildford</p>{FILLER}", "")[0], "confirmed")

    def test_name_alone_is_only_possible(self):
        # Could be a Smith Roofing in Leeds.
        verdict, why = fp.judge_site(company("SMITH ROOFING LTD"), f"<h1>Smith Roofing</h1><p>Leeds LS1 1AA</p>{FILLER}", "")
        self.assertEqual(verdict, "possible")

    def test_initials_can_be_dropped(self):
        item = company("J SMITH ROOFING LTD")
        self.assertEqual(fp.judge_site(item, f"<h1>Smith Roofing</h1><p>Guildford</p>{FILLER}", "")[0], "confirmed")

    def test_parked_and_unrelated_sites(self):
        item = company("SMITH ROOFING LTD")
        self.assertEqual(fp.judge_site(item, f"<h1>This domain is for sale</h1>{FILLER}", "")[0], "no")
        self.assertEqual(fp.judge_site(item, f"<h1>Jones Plumbing</h1>{FILLER}", "")[0], "no")


class FindWebsite(unittest.TestCase):
    def test_takes_the_first_confirmed_and_skips_directories(self):
        pages = {
            "https://smithroofing.co.uk/": (f"<h1>Smith Roofing</h1><p>Leeds</p>{FILLER}", "https://smithroofing.co.uk/"),
            "https://smithroofing.com/": ("<p>x</p>", "https://www.checkatrade.com/trades/smith"),
            "https://smith-roofing.co.uk/": (f"<h1>Smith Roofing</h1><p>Guildford</p>{FILLER}", "https://www.smith-roofing.co.uk/"),
        }
        fetch = lambda url: pages.get(url, (None, None))  # noqa: E731
        live = lambda d: f"https://{d}/" in pages  # noqa: E731
        site = fp.find_website(company("SMITH ROOFING LTD"), "Guildford", fetch, live)
        self.assertEqual((site["status"], site["website"]), ("confirmed", "smith-roofing.co.uk"))

    def test_falls_back_to_possible_then_none(self):
        pages = {"https://smithroofing.co.uk/": (f"<h1>Smith Roofing</h1><p>Leeds</p>{FILLER}", "https://smithroofing.co.uk/")}
        fetch = lambda url: pages.get(url, (None, None))  # noqa: E731
        site = fp.find_website(company("SMITH ROOFING LTD"), "Guildford", fetch, lambda d: f"https://{d}/" in pages)
        self.assertEqual(site["status"], "possible")
        site = fp.find_website(company("SMITH ROOFING LTD"), "Guildford", lambda u: (None, None), lambda d: False)
        self.assertEqual(site["status"], "none")


class Contacts(unittest.TestCase):
    def test_email_prefers_their_own_domain_and_inbox(self):
        page = (
            '<a href="mailto:john.smith@gmail.com">x</a> sales@smithroofing.co.uk info@smithroofing.co.uk '
            "logo@2x.png support@wix.com orders@supplier.co.uk"
        )
        self.assertEqual(fp.find_email(page, "smithroofing.co.uk"), "info@smithroofing.co.uk")
        self.assertEqual(fp.find_email('<a href="mailto:bob@gmail.com">', "smithroofing.co.uk"), "bob@gmail.com")
        self.assertEqual(fp.find_email("orders@supplier.co.uk", "smithroofing.co.uk"), "")

    def test_cloudflare_protected_email(self):
        key = 0x42
        hexed = f"{key:02x}" + "".join(f"{ord(c) ^ key:02x}" for c in "info@smith.co.uk")
        page = f'<a href="/cdn-cgi/l/email-protection#{hexed}">[email&#160;protected]</a>'
        self.assertEqual(fp.find_email(page, "smith.co.uk"), "info@smith.co.uk")

    def test_phone(self):
        self.assertEqual(fp.find_phone('<a href="tel:+441483123456">Call</a>'), "01483 123456")
        self.assertEqual(fp.find_phone("<p>Call 07700 900123 today</p>"), "07700 900123")
        self.assertEqual(fp.find_phone("<p>Call 020 7946 0000</p>"), "020 7946 0000")
        self.assertEqual(fp.find_phone("<p>No number, company 12345678</p>"), "")


class EndToEnd(unittest.TestCase):
    """The whole run against a pretend register and pretend websites."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.outreach = Path(self.tmp.name) / "outreach"
        self.outreach.mkdir()
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Outreach"
        ws.append(["Business", "Website", "Status"])
        ws.append(["Already Known Roofing", "knownroof.co.uk", "Lost / not interested"])
        wb.save(self.outreach / "outreach-master.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def run_finder(self, *extra):
        register = [
            company("SMITH ROOFING LTD", "00000001"),
            company("JONES ROOFING LTD", "00000002"),
            company("NO WEB ROOFING LTD", "00000003"),
            company("ALREADY KNOWN ROOFING LTD", "00000004"),
            company("ACME HOLDINGS LTD", "00000005"),
            company("SURREY LOFTS LTD", "00000006", sic=("41202",)),
            {**company("PARTNERSHIP ROOFERS", "00000007"), "company_type": "limited-partnership"},
        ]
        sites = {
            "https://smithroofing.co.uk/": (
                f'<h1>Smith Roofing</h1><p>Guildford GU1 2AB</p><a href="tel:01483111222">Call</a>'
                f'<a href="/contact">Contact</a>{FILLER}',
                "https://smithroofing.co.uk/",
            ),
            "https://smithroofing.co.uk/contact": ("<p>Email info@smithroofing.co.uk</p>", "https://smithroofing.co.uk/contact"),
            "https://jonesroofing.co.uk/": (f"<h1>Jones Roofing</h1><p>Leeds</p>{FILLER}", "https://jonesroofing.co.uk/"),
        }

        def fake_get(path, key):
            if "/officers" in path:
                return {"items": [{"name": "SMITH, John", "officer_role": "director", "appointed_on": "2015-01-01"}]}
            if "sic_codes=43910" in path:
                return {"items": register, "hits": len(register)}
            return {"items": [], "hits": 0}

        buf = io.StringIO()
        with mock.patch.object(fp, "_get", fake_get), mock.patch.object(fp, "outreach_dir", lambda: self.outreach), \
                mock.patch.object(fp, "resolves", lambda d: f"https://{d}/" in sites), \
                mock.patch("site_teardown.fetch_html", lambda url, timeout=20: sites.get(url, (None, None))), \
                mock.patch.dict(os.environ, {"COMPANIES_HOUSE_API_KEY": "1a2b3c4d-5e6f-7a8b-9c0d-1e2f3a4b5c6d"}), \
                redirect_stdout(buf):
            fp.main(["--trades", "roofing", "--areas", "Guildford", *extra])
        return buf.getvalue()

    def test_builds_a_sheet_with_three_tabs(self):
        import openpyxl

        out = self.run_finder()
        files = [p for p in self.outreach.glob("roofing-guildford-*.xlsx")]
        self.assertEqual(len(files), 1, out)
        wb = openpyxl.load_workbook(files[0])
        self.assertEqual(wb.sheetnames, ["Outreach", "Check website", "No website"])

        def rows(tab):
            ws = wb[tab]
            header = [c.value for c in ws[1]]
            return [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

        outreach = rows("Outreach")
        self.assertEqual([r["Business"] for r in outreach], ["Smith Roofing"])
        smith = outreach[0]
        self.assertEqual(smith["Website"], "smithroofing.co.uk")
        self.assertEqual(smith["Email"], "info@smithroofing.co.uk")
        self.assertEqual(smith["Phone"], "01483 111222")
        self.assertEqual(smith["Contact name"], "John Smith")
        self.assertEqual(smith["Company type"], "Ltd")
        self.assertEqual(smith["Trade"], "Roofing")
        self.assertEqual(smith["Status"], "New")
        self.assertEqual([r["Possible website"] for r in rows("Check website")], ["jonesroofing.co.uk"])
        self.assertEqual([r["Business"] for r in rows("No website")], ["No Web Roofing"])
        # Known firm, holding company, partnership and non-roofer all left out.
        self.assertIn("1 already in your sheets", out)

    def test_email_only_keeps_just_firms_with_an_email(self):
        import openpyxl

        out = self.run_finder("--email-only")
        wb = openpyxl.load_workbook(next(self.outreach.glob("roofing-guildford-*.xlsx")))
        names = [r[0].value for tab in wb.worksheets for r in tab.iter_rows(min_row=2)]
        self.assertEqual(names, ["Smith Roofing"])
        self.assertIn("left out for having no email", out)
        self.assertIn("Only 1 of the 100", out)

    def test_max_counts_kept_firms_when_filtering(self):
        import openpyxl

        self.run_finder("--email-only", "--max", "1")
        wb = openpyxl.load_workbook(next(self.outreach.glob("roofing-guildford-*.xlsx")))
        self.assertEqual(sum(ws.max_row - 1 for ws in wb.worksheets), 1)

    def test_email_only_needs_the_website_search(self):
        with self.assertRaises(SystemExit) as e:
            self.run_finder("--email-only", "--no-websites")
        self.assertIn("website search", str(e.exception))

    def test_count_only_writes_nothing(self):
        out = self.run_finder("--count-only")
        self.assertIn("3 new to you", out)
        self.assertEqual(list(self.outreach.glob("roofing-*.xlsx")), [])

    def test_the_output_is_a_sheet_push_prospects_reads(self):
        self.run_finder()
        path = next(self.outreach.glob("roofing-guildford-*.xlsx"))
        buf = io.StringIO()
        env = {"PROSPECTS_API_SECRET": "x" * 40}
        with mock.patch.dict(os.environ, env), redirect_stdout(buf), mock.patch("sys.argv", ["p", "--sheet", str(path), "--dry-run"]):
            import push_prospects

            push_prospects.main()
        self.assertIn("1 prospects (1 email, 0 letter)", buf.getvalue())


if __name__ == "__main__":
    unittest.main()
