"""Run the panel from your phone: Telegram buttons for the day-to-day loop.

Runs inside the panel (so it shares its one-job-at-a-time lock and every safeguard), whenever the
panel is open and TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID are set. Turn on "Start the panel with
Windows" on the Today tab and it's always there while the PC is on.

It only ever answers YOUR chat (TELEGRAM_CHAT_ID): anyone else who finds the bot gets silence. It
asks Telegram for new taps (outbound only - no port is opened on your PC, nothing to configure).

    📊 Status      what's running, the autopilot, today's sends, replies and calls waiting, red flags
    💬 Replies     each reply waiting: ✨ AI draft -> Send / Edit / Discard, call back, not interested
    📞 Calls       the hottest preview viewers (tap the number to ring): outcome buttons, 🔎 research
    ✉️ Emails      waiting to email: make a batch, preview it, send it (asks first), follow-ups
    🗂 Lists       run the whole list on any list with checks still to do
    🤖 Autopilot   run it now, or stop a run
    📈 Scorecard   the numbers
    📜 Log / ⏹ Stop   the running job's last lines, or stop it

Each morning from 8:30 it texts you the day's plan (and on Sundays the week against the last), then
anyone who opened their preview twice or more without a reply
or a call, with a follow-up written for them - ✅ Send it, ✏️ Edit or 🗑 Discard (nudges.py) - and
anyone who opened their quote two days ago without accepting it (quote_chase.py), and each client a
month after launch, with a review and referral ask ready to paste (daily.py). Type "objections" for
answers to what people say on calls.

Anything started here texts you its result when it finishes. Kept on the panel only (screen work, or
too risky for a stray tap): finding new firms, building and publishing client sites, Settings, and
erasing someone's data.
"""

from __future__ import annotations

import hashlib
import json
import threading
import time
from datetime import datetime
import urllib.error
import urllib.request

MENU = [["📊 Status", "💬 Replies", "📞 Calls"], ["✉️ Emails", "🗂 Lists", "🤖 Autopilot"], ["📈 Scorecard", "📜 Log", "⏹ Stop"]]
CALLS_SHOWN = 6
LOG_LINES = 15
SEND_GAP_MIN = 5
BATCH_SIZE = 20


def short(key: str) -> str:
    return hashlib.sha1(key.encode()).hexdigest()[:10]


def mobile(phone: str) -> str:
    """A UK mobile as WhatsApp wants it (447...), or "" for a landline or anything unclear."""
    digits = "".join(c for c in phone or "" if c.isdigit())
    if digits.startswith("07") and len(digits) == 11:
        return "44" + digits[1:]
    if digits.startswith("447") and len(digits) == 12:
        return digits
    return ""


def clip(text: str, n: int = 3800) -> str:
    return text if len(text) <= n else text[: n - 20] + "\n... (cut short)"


class Bot:
    def __init__(self, panel, token: str, chat: str, http=None) -> None:
        self.panel, self.token, self.chat = panel, token, str(chat).strip()
        self.http = http or self._http
        self.items: dict[str, dict] = {}  # short id -> {"kind", "item", "sheet"} from the last listing
        self.drafts: dict[str, str] = {}  # short id -> reply text waiting for Send
        self.awaiting: str = ""  # short id whose reply you're typing
        self.lists: list[str] = []
        self.offset = 0
        self.stop_event = threading.Event()

    # ------------------------------------------------------------ Telegram
    def _http(self, method: str, payload: dict, timeout: float = 30) -> dict:
        req = urllib.request.Request(f"https://api.telegram.org/bot{self.token}/{method}",
                                     data=json.dumps(payload).encode(), headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=timeout) as res:
            return json.loads(res.read())

    def api(self, method: str, **payload) -> dict:
        try:
            return self.http(method, payload)
        except (urllib.error.URLError, TimeoutError, OSError, ValueError):
            return {}

    def send(self, text: str, buttons: list[list[tuple[str, str]]] | None = None, menu: bool = False) -> None:
        payload = {"chat_id": self.chat, "text": clip(text), "disable_web_page_preview": True}
        if buttons:
            payload["reply_markup"] = {"inline_keyboard": [[{"text": t, "url": d} if d.startswith("https://") else {"text": t, "callback_data": d}
                                                             for t, d in row] for row in buttons]}
        elif menu:
            payload["reply_markup"] = {"keyboard": [[{"text": t} for t in row] for row in MENU], "resize_keyboard": True,
                                       "is_persistent": True}
        self.api("sendMessage", **payload)

    def mine(self, chat_id) -> bool:
        return bool(self.chat) and str(chat_id) == self.chat

    # ------------------------------------------------------------ loop
    def run_forever(self) -> None:
        wait = 5
        while not self.stop_event.is_set():
            try:
                res = self.http("getUpdates", {"offset": self.offset, "timeout": 50,
                                               "allowed_updates": ["message", "callback_query"]}, 70)
                wait = 5
            except (urllib.error.URLError, TimeoutError, OSError, ValueError):
                self.stop_event.wait(wait)  # offline, or the PC just woke: try again, slower each time
                wait = min(wait * 2, 300)
                continue
            for update in res.get("result") or []:
                self.offset = max(self.offset, int(update.get("update_id", 0)) + 1)
                try:
                    self.handle(update)
                except Exception as e:  # noqa: BLE001 - one bad tap never stops the bot
                    self.send(f"That didn't work: {type(e).__name__}: {e}")
            self.maybe_morning()

    def maybe_morning(self, now_local: datetime | None = None) -> bool:
        """Once a day from 8:30: the morning nudges, in the background so taps still answer."""
        import nudges

        day_file = self.panel.OUTREACH / "nudge-day.txt"
        try:
            last = day_file.read_text(encoding="utf-8").strip()
        except OSError:
            last = ""
        now_local = now_local or datetime.now()
        if not nudges.due(now_local, last):
            return False
        try:
            day_file.write_text(now_local.date().isoformat(), encoding="utf-8")  # first, so a crash never repeats it
        except OSError:
            return False
        threading.Thread(target=self.morning, daemon=True).start()
        return True

    def morning(self) -> None:
        """8:30: the day's plan, the week on Sundays, then each nudge, quote chaser and referral ask."""
        import email_batches
        import nudges

        p = self.panel
        data: dict = {}
        quotes: list = []
        try:
            data = self._all_calls()
            quotes = self._open_quotes()
            import daily

            today = datetime.now().date()
            goals = p.load_settings().get("WEEKLY_TARGETS") or ""
            plan = daily.plan({} if data.get("message") else data, len(quotes), email_batches.remaining(p.OUTREACH, None))
            # Monday starts a fresh week: the count shows from Tuesday, and in full on Sunday.
            self.send(plan + ("\n\n" + daily.progress(p.OUTREACH, today, goals) if today.weekday() not in (0, 6) else ""))
            if today.weekday() == 6:
                self.send(daily.week(p.OUTREACH, email_batches.preview_views(), today) + "\n" + daily.progress(p.OUTREACH, today, goals))
        except Exception as e:  # noqa: BLE001 - a bad morning never stops the bot
            self.send(f"Today's plan didn't work: {type(e).__name__}: {e}")
        try:
            items = [] if not data or data.get("message") else nudges.candidates(data, p.OUTREACH)
            name = p.load_settings().get("MAIL_FROM_NAME") or ""
            for item in items:
                d = nudges.draft(item, name)
                sid = self._remember("nudge", {**item, "subject": d["subject"], "message_id": d["message_id"]})
                self.drafts[sid] = d["text"]
                nudges.mark(p.OUTREACH, item["website"], item["business"], datetime.now().date())
                ring = f"Best move: ring {item['phone']}." if item.get("phone") else "No phone on your list."
                self.send(f"☀️ {item['business']}{' - ' + item['contact'] if item.get('contact') else ''} opened their preview "
                          f"{item['views']} times ({item.get('seconds', 0)}s on it) and hasn't replied.\n{ring} Or send this:")
                self.offer_draft(sid)
        except Exception as e:  # noqa: BLE001
            self.send(f"Morning nudges didn't work: {type(e).__name__}: {e}")
        try:
            self.chase_quotes(quotes)
        except Exception as e:  # noqa: BLE001
            self.send(f"Quote chasers didn't work: {type(e).__name__}: {e}")
        try:
            self.ask_referrals()
        except Exception as e:  # noqa: BLE001
            self.send(f"Referral reminders didn't work: {type(e).__name__}: {e}")

    def _open_quotes(self) -> list:
        import os

        import quote_chase

        settings = self.panel.load_settings()
        secret = settings.get("PROSPECTS_API_SECRET") or ""
        if len(secret) < 32:
            return []
        api = (settings.get("DASHBOARD_API_URL") or "https://admin.scalardigital.co.uk").rstrip("/")
        tenant = os.environ.get("SCALAR_TENANT_ID", "abdc6408-1fd5-4fb6-9c4c-53600b571a6d")
        return quote_chase.fetch(api, secret, tenant)

    def ask_referrals(self) -> None:
        """30 days after a launch: ask for a review and a referral - the words ready to paste (daily.py)."""
        import daily

        p = self.panel
        settings = p.load_settings()
        site = (settings.get("SITE_URL") or "https://www.scalardigital.co.uk").rstrip("/")
        for row in daily.referrals_due(p.OUTREACH, datetime.now().date()):
            daily.mark_asked(p.OUTREACH, row["new_site"], datetime.now().date())
            self.send(f"🤝 {row.get('business') or row['new_site']} went live a month ago - ask for a "
                      f"{'review and a ' if settings.get('REVIEW_LINK') else ''}referral. Send them this (WhatsApp or email):")
            self.send(daily.referral_text(row, site, settings.get("REVIEW_LINK") or ""))

    def chase_quotes(self, quotes: list | None = None) -> None:
        """Quotes opened but not accepted after two days: ring them, or send the follow-up (quote_chase.py)."""
        import nudges
        import quote_chase

        p = self.panel
        settings = p.load_settings()
        quotes = quote_chase.due(self._open_quotes() if quotes is None else quotes, p.OUTREACH)
        sent = nudges.sent_rows(p.OUTREACH)
        for q in quotes:
            firm = self.find_by_slug(q.get("slug") or "") or {}
            quote_chase.mark(p.OUTREACH, q, datetime.now().date())
            name = q.get("business_name") or firm.get("business") or "A prospect"
            ring = f"Ring {firm['phone']} - best move." if firm.get("phone") else "No phone on your list."
            self.send(f"💷 {name} opened their {quote_chase.pounds(q.get('total_pence'))} quote "
                      f"{q.get('view_count')} time(s) and hasn't accepted.\n{ring}\n{q.get('quote_url', '')}")
            if not (firm.get("email") and firm.get("sheet") and firm.get("key")):
                continue
            first = sent.get(firm["email"].strip().lower()) or {}
            sid = self._remember("quote", {**firm, "subject": first.get("subject") or "", "message_id": first.get("message_id") or ""})
            self.drafts[sid] = quote_chase.draft(q, firm, settings.get("MAIL_FROM_NAME") or "")
            self.offer_draft(sid)

    def handle(self, update: dict) -> None:
        if "callback_query" in update:
            q = update["callback_query"]
            if not self.mine(((q.get("message") or {}).get("chat") or {}).get("id")):
                return
            self.api("answerCallbackQuery", callback_query_id=q.get("id"))
            data = str(q.get("data") or "")
            if data == "inbound":  # the website's "asked for their report" alert - its own text is the lead
                self.inbound(str((q.get("message") or {}).get("text") or ""))
            elif data.startswith("callnow:"):  # the website's "just opened their preview" alert
                self.call_now(data.partition(":")[2])
            else:
                self.on_button(data)
            return
        msg = update.get("message") or {}
        if not self.mine((msg.get("chat") or {}).get("id")):
            return
        text = (msg.get("text") or "").strip()
        if not text:
            return
        handler = {"📊 Status": self.status, "💬 Replies": self.replies, "📞 Calls": self.calls, "✉️ Emails": self.emails,
                   "🗂 Lists": self.list_menu, "🤖 Autopilot": self.autopilot, "📈 Scorecard": self.scorecard,
                   "📜 Log": self.log, "⏹ Stop": self.stop_job}.get(text)
        if handler:
            self.awaiting = ""
            handler()
        elif text.lower() in ("objections", "/objections"):
            import objections

            self.send(objections.text())
        elif text.lower() in ("/start", "/help", "menu", "help"):
            self.awaiting = ""
            self.send("Scalar panel - pick from the buttons below. Type \"objections\" for answers to what people say on calls. "
                      "Only this chat can use it.", menu=True)
        elif self.awaiting:
            sid, self.awaiting = self.awaiting, ""
            self.drafts[sid] = text
            self.offer_draft(sid)
        else:
            self.send("Pick from the buttons below.", menu=True)

    # ------------------------------------------------------------ screens
    def status(self) -> None:
        import autopilot
        import email_batches
        import send_email
        import sending_health
        from datetime import date

        p = self.panel
        settings = p.load_settings()
        lines = []
        job = p.JOB.state()
        if job["running"]:
            mins = int((time.time() - job["started"]) // 60)
            last = next((ln for ln in reversed(job["lines"]) if ln.strip() and not ln.startswith(">")), "")
            lines.append(f"Running: {job['label']} ({mins} min) - {last[:120]}")
        if autopilot.is_running():
            now = autopilot.now_doing()
            lines.append(f"Autopilot running ({now.get('minutes', 0)} min): {now.get('last', '')[:120]}")
        else:
            cfg = autopilot.load_config()
            lines.append(f"Autopilot: on, daily at {cfg.get('time')}" if cfg.get("enabled") else "Autopilot: off")
        import mail_accounts

        inboxes = [a.address for a in mail_accounts.accounts({**settings})]
        today = send_email.today_summary(p.OUTREACH, date.today(), inboxes)
        lines.append(f"Sent today: {today['sent']} of {today['cap']}")
        lines.append(f"Waiting to email: {email_batches.remaining(p.OUTREACH, None)}")
        replies = self._replies_waiting()
        lines.append(f"Replies waiting: {len(replies)}")
        stop = sending_health.stop_sending(p.OUTREACH)
        if stop:
            lines.append("⚠️ " + stop)
        bad = [h for h in p.health(p.OUTREACH, settings) if h.get("ok") is False]
        lines += [f"⚠️ {h['name']}: {h['text']}" for h in bad]
        self.send("\n".join(lines), menu=True)

    def _replies_waiting(self) -> list[dict]:
        import csv

        path = self.panel.OUTREACH / "replies.csv"
        if not path.exists():
            return []
        with path.open(encoding="utf-8") as f:
            return [r for r in csv.DictReader(f) if r.get("kind") in ("interested", "read it") and not r.get("handled")]

    def _all_calls(self) -> dict:
        """Calls from every list, the dashboard asked once - newest list first, each firm once."""
        import calls

        p = self.panel
        settings = p.load_settings()
        secret = settings["PROSPECTS_API_SECRET"]
        if len(secret) < 32:
            return {"message": "Add PROSPECTS_API_SECRET in the panel's Settings first."}
        api = (settings["DASHBOARD_API_URL"] or "https://admin.scalardigital.co.uk").rstrip("/")
        site = (settings["SITE_URL"] or "https://www.scalardigital.co.uk").rstrip("/")
        import os

        tenant = os.environ.get("SCALAR_TENANT_ID", "abdc6408-1fd5-4fb6-9c4c-53600b571a6d")
        try:
            activity = calls.fetch_activity(api, secret, tenant)
        except calls.DashboardMissing as e:
            return {"message": str(e)}
        out: dict[str, list] = {"callbacks": [], "viewing": [], "replied": []}
        seen: set[tuple[str, str]] = set()
        names = sorted(p.sheets(), key=lambda n: (p.OUTREACH / n).stat().st_mtime, reverse=True)
        for name in names:
            try:
                got = calls.call_list(p.OUTREACH, name, secret, site, activity)
            except Exception:  # noqa: BLE001 - a sheet open in Excel: skip it
                continue
            for section in out:
                for item in got.get(section) or []:
                    # A reply shows on every list; the list that has the firm knows its contact and phone.
                    mark = (section, item.get("website") or item.get("key"))
                    if mark in seen:
                        continue
                    if section == "replied" and not (item.get("contact") or item.get("phone")) and name != names[-1]:
                        continue
                    seen.add(mark)
                    out[section].append({**item, "sheet": item.get("sheet") or name})
        out["viewing"].sort(key=lambda i: -i.get("heat", 0))
        return out

    def _script(self, item: dict) -> str:
        import call_script

        return "\n".join(call_script.script(item, self.panel.OUTREACH, self.panel.load_settings().get("MAIL_FROM_NAME") or ""))

    def after_no_answer(self, item: dict) -> None:
        """No answer: a ready message they'll see straight away - one tap opens WhatsApp with it typed."""
        import urllib.parse

        from saved_replies import first_name

        me = (self.panel.load_settings().get("MAIL_FROM_NAME") or "").split()
        name = first_name(item.get("contact", ""))
        link = (item.get("preview") or "").replace("src=dashboard", "src=whatsapp")
        text = (f"Hi{' ' + name if name != 'there' else ''}, it's {me[0] if me else 'Scalar Digital'}"
                f"{' from Scalar Digital' if me else ''} - just tried to ring you. Here's the page I put together for "
                f"{item['business']}: {link} - happy to talk it through whenever suits.")
        number = mobile(item.get("phone", ""))
        if number:
            self.send(f"Send them this on WhatsApp?\n\n{text}",
                      [[("💬 Open WhatsApp with it", f"https://wa.me/{number}?{urllib.parse.urlencode({'text': text})}")]])
        elif link:
            self.send(f"Their number isn't a mobile, so no WhatsApp. If you have one for them, send:\n\n{text}")

    def _remember(self, kind: str, item: dict) -> str:
        sid = short(f"{kind}|{item.get('sheet')}|{item.get('key')}")
        self.items[sid] = {"kind": kind, "item": item}
        return sid

    def replies(self) -> None:
        data = self._all_calls()
        if data.get("message"):
            self.send(data["message"])
            return
        items = data["replied"]
        if not items:
            self.send("No replies waiting. 👍", menu=True)
            return
        for i in items[:8]:
            sid = self._remember("reply", i)
            phone = f"\n📞 {i['phone']}" if i.get("phone") else ""
            self.send(f"💬 {i['business']}{' - ' + i['contact'] if i.get('contact') else ''}{phone}\n{i.get('email', '')}\n\n\"{i.get('snippet', '')[:600]}\"",
                      [[("✨ AI draft", f"ai:{sid}"), ("✏️ Write reply", f"wr:{sid}")],
                       [("📞 Call back tomorrow", f"cb:{sid}"), ("🚫 Not interested", f"ni?:{sid}")],
                       [("✓ Dealt with", f"done:{sid}")]])

    def calls(self) -> None:
        data = self._all_calls()
        if data.get("message"):
            self.send(data["message"])
            return
        items = data["callbacks"] + data["viewing"]
        if not items:
            self.send("Nobody new has opened their preview since your last calls.", menu=True)
            return
        for i in items[:CALLS_SHOWN]:
            sid = self._remember("call", i)
            why = (f"Call back {'(' + str(i['overdue']) + ' days overdue)' if i.get('overdue') else 'today'}"
                   + (f' - "{i["note"]}"' if i.get("note") else "")) if "due" in i else \
                f"{i.get('views', 0)} visit(s), {i.get('seconds', 0)}s on the page"
            self.send(f"📞 {i['business']}{' - ' + i['contact'] if i.get('contact') else ''}\n"
                      f"{i.get('phone') or 'no phone - check their site'}\n{why}\n{i.get('preview', '')}\n\n{self._script(i)}",
                      [[("No answer", f"o:{sid}:No answer"), ("Call back tmrw", f"cb:{sid}")],
                       [("👍 Interested", f"o:{sid}:Interested"), ("🚫 Not interested", f"ni?:{sid}")],
                       [("🏆 Won", f"o:{sid}:Won"), ("🔎 Research", f"rs:{sid}")]])
        if len(items) > CALLS_SHOWN:
            self.send(f"...and {len(items) - CALLS_SHOWN} more on the panel's Calls tab.")

    def emails(self) -> None:
        import email_batches

        p = self.panel
        waiting = email_batches.remaining(p.OUTREACH, None)
        pending = email_batches._rows(p.OUTREACH / email_batches.PENDING)[1]
        lines = [f"{waiting} waiting to be emailed."]
        if pending:
            lines.append(f"Batch ready: {len(pending)} emails ({pending[0].get('batch', '')}).")
        self.send("\n".join(lines), [[(f"Make batch of {BATCH_SIZE}", "run:batch"), ("👀 Preview", "preview")],
                                     [(f"🚀 Send batch ({SEND_GAP_MIN} min apart)", "send?")],
                                     [("Make follow-ups", "run:followups"), ("Send follow-ups", "sendf?")]])

    def list_menu(self) -> None:
        p = self.panel
        names = sorted(p.sheets(), key=lambda n: (p.OUTREACH / n).stat().st_mtime, reverse=True)[:8]
        self.lists = names
        rows = []
        for i, n in enumerate(names):
            left = (p.progress(n) or {}).get("remaining") or 0
            rows.append([(f"Run {n.removesuffix('.xlsx')}" + (f" ({left} to check)" if left else ""), f"list:{i}")])
        self.send("Run the whole list on:" if rows else "No lists in outreach/ yet.", rows or None)

    def autopilot(self) -> None:
        import autopilot

        if autopilot.is_running():
            now = autopilot.now_doing()
            self.send(f"Autopilot running ({now.get('minutes', 0)} min): {now.get('last', '')}", [[("⏹ Stop the autopilot", "apstop?")]])
        else:
            cfg = autopilot.load_config()
            self.send(f"Autopilot is {'on, daily at ' + cfg.get('time', '') if cfg.get('enabled') else 'off'}.", [[("▶️ Run it now", "run:autopilot_now")]])

    def scorecard(self) -> None:
        import email_batches as eb
        import scorecard

        self.send("\n".join(scorecard.scorecard(self.panel.OUTREACH, eb.preview_views())))

    def log(self) -> None:
        job = self.panel.JOB.state()
        tail = [ln for ln in job["lines"] if ln.strip()][-LOG_LINES:]
        self.send("\n".join(tail) if tail else "Nothing has run since the panel started.")

    def stop_job(self) -> None:
        if not self.panel.JOB.running():
            self.send("Nothing is running." + (" (The autopilot is - use 🤖 Autopilot to stop it.)" if self._autopilot_on() else ""))
            return
        self.panel.JOB.stop()
        self.send("Stopping. Anything already sent stays sent.")

    def _autopilot_on(self) -> bool:
        import autopilot

        return autopilot.is_running()

    # ------------------------------------------------------------ buttons
    def on_button(self, data: str) -> None:
        verb, _, rest = data.partition(":")
        if verb == "run":
            self.start(rest)
        elif verb == "list":
            i = int(rest) if rest.isdigit() else -1
            if 0 <= i < len(self.lists):
                self.start("all", {"sheet": self.lists[i]})
            else:
                self.send("That list button is out of date - tap 🗂 Lists again.")
        elif verb == "preview":
            self.preview()
        elif verb == "send?":
            self.send(f"Send today's batch now, {SEND_GAP_MIN} min apart?", [[("Yes, send", "run:send_batch"), ("Cancel", "cancel")]])
        elif verb == "sendf?":
            self.send("Send the follow-ups now?", [[("Yes, send", "run:send_followups"), ("Cancel", "cancel")]])
        elif verb == "apstop?":
            self.send("Stop the autopilot run? Anything already sent stays sent.", [[("Yes, stop it", "apstop"), ("Cancel", "cancel")]])
        elif verb == "apstop":
            import autopilot

            self.send(autopilot.stop())
        elif verb == "cancel":
            self.send("OK - nothing done.")
        else:
            self.on_item(verb, rest)

    def on_item(self, verb: str, rest: str) -> None:
        sid, _, extra = rest.partition(":")
        entry = self.items.get(sid)
        if not entry:
            self.send("That button is out of date (the panel restarted) - tap 💬 Replies or 📞 Calls again.")
            return
        item = entry["item"]
        body = {"sheet": item.get("sheet"), "key": item.get("key"), "business": item.get("business"),
                "website": item.get("website"), "email": item.get("email")}
        p = self.panel
        if verb == "o":
            out, _ = p.call_action({**body, "outcome": extra})
            self.send(out.get("message") or out.get("error", ""))
            if extra == "No answer" and not out.get("error"):
                self.after_no_answer(item)
        elif verb == "cb":
            out, _ = p.call_action({**body, "outcome": "Call back", "due": "tomorrow"})
            self.send(out.get("message") or out.get("error", ""))
        elif verb == "ni?":
            import calls

            reasons = [(r, f"ni:{sid}:{r}") for r in calls.LOST_REASONS]
            self.send(f"{item['business']}: not interested - they won't be contacted again. Why?",
                      [reasons[:3], [*reasons[3:], ("Cancel", "cancel")]])
        elif verb == "ni":
            out, _ = p.call_action({**body, "outcome": "Not interested", "reason": extra})
            self.send(out.get("message") or out.get("error", ""))
        elif verb == "done":
            import calls

            calls.mark_reply_handled(p.OUTREACH, item["business"])
            self.send(f"{item['business']}: marked as dealt with.")
        elif verb == "rs":
            out, _ = p.research_action({"business": item["business"], "website": item["website"], "area": item.get("area", "")})
            self.send("\n".join(out.get("lines") or [out.get("error", "")]) + (f"\nReviews: {out['google']}" if out.get("google") else ""))
        elif verb == "ai":
            self.send("Writing a draft...")
            out, _ = p.reply_ai_action({k: item.get(k, "") for k in ("business", "trade", "area", "website", "mobile_score", "subject", "message",
                                                                    "preview", "contact")})
            if out.get("error"):
                self.send(out["error"])
                return
            self.drafts[sid] = out["text"]
            self.offer_draft(sid)
        elif verb == "wr":
            self.awaiting = sid
            self.send(f"Type your reply to {item['business']} as your next message - you'll see it before it goes.")
        elif verb == "edit":
            self.awaiting = sid
            self.send("Send the reply as you want it (copy the draft above, change it, send it here).")
        elif verb == "discard":
            self.drafts.pop(sid, None)
            self.send("Draft thrown away.")
        elif verb == "sendreply":
            text = self.drafts.pop(sid, "")
            if not text:
                self.send("That draft has gone - make it again.")
                return
            out, _ = p.reply_send_action({**body, "to": item.get("email", ""), "text": text, "message_id": item.get("message_id", ""),
                                          "subject": item.get("subject", ""), "name": "Telegram"})
            self.send(out.get("message") or out.get("error", ""))

    def find_by_slug(self, slug: str) -> dict | None:
        """The firm whose preview this is, from your lists - with the phone only this PC has."""
        import csv

        import openpyxl

        import calls
        import overrides
        from push_prospects import domain_of, make_slug

        p = self.panel
        settings = p.load_settings()
        secret = settings["PROSPECTS_API_SECRET"]
        found = {}
        contacts = p.OUTREACH / "contacts-found.csv"
        if contacts.exists():
            with contacts.open(encoding="utf-8") as f:
                found = {r["website"]: r for r in csv.DictReader(f) if r.get("website")}
        site = (settings.get("SITE_URL") or "https://www.scalardigital.co.uk").rstrip("/")
        for name in p.sheets():
            try:
                wb = openpyxl.load_workbook(p.OUTREACH / name, read_only=True, data_only=True)
            except Exception:  # noqa: BLE001 - open in Excel: try the others
                continue
            try:
                for tab in ("Outreach", "Check website"):
                    if tab not in wb.sheetnames:
                        continue
                    rows = wb[tab].iter_rows(values_only=True)
                    header = [str(h).strip() if h else "" for h in next(rows, ())]
                    for values in rows:
                        row = dict(zip(header, values))
                        business = str(row.get("Business") or "").strip()
                        domain = domain_of(str(row.get("Website") or row.get("Possible website") or "")) or ""
                        if not business or not domain or make_slug(business, domain, secret) != slug:
                            continue
                        extra = found.get(domain) or {}
                        phone = next((str(row[c]).strip() for c in calls.PHONE_COLUMNS if str(row.get(c) or "").strip()), "") or extra.get("phone", "")
                        return {"key": overrides.row_key(row), "sheet": name, "business": business, "website": domain,
                                "contact": str(row.get("Contact name") or "").strip() or extra.get("contact", ""),
                                "email": str(row.get("Email") or "").strip() or extra.get("email", ""), "phone": phone,
                                "area": str(row.get("Area") or row.get("Town") or "").strip(),
                                "preview": f"{site}/for/{slug}?src=dashboard"}
            finally:
                wb.close()
        return None

    def call_now(self, slug: str) -> None:
        item = self.find_by_slug(slug)
        if not item:
            self.send("Couldn't find that firm in your lists on this PC - check Calls on the panel.")
            return
        sid = self._remember("call", item)
        who = f" - ask for {item['contact']}" if item.get("contact") else ""
        self.send(f"📞 {item['business']}{who}\n{item['phone'] or 'No phone on your list - check their site: ' + item['website']}\n\n"
                  f"{self._script(item)}",
                  [[("No answer", f"o:{sid}:No answer"), ("Call back tmrw", f"cb:{sid}")],
                   [("👍 Interested", f"o:{sid}:Interested"), ("🚫 Not interested", f"ni?:{sid}")]])
        out, _ = self.panel.research_action({"business": item["business"], "website": item["website"], "area": item.get("area", "")})
        if out.get("lines"):
            self.send("\n".join(out["lines"]))

    def inbound(self, text: str) -> None:
        """"Add to pipeline" on a website lead: into inbound.xlsx, preview built, then a personal reply offered."""
        import inbound
        import overrides
        from push_prospects import domain_of, make_slug

        lead = inbound.parse(text)
        if not lead:
            self.send("That alert doesn't have what's needed (business, email, website) - add them by hand.")
            return
        p = self.panel
        status, message = inbound.add(p.OUTREACH, lead)
        self.send(message)
        if status not in ("added", "already"):
            return
        settings = p.load_settings()
        site = (settings.get("SITE_URL") or "https://www.scalardigital.co.uk").rstrip("/")
        domain = domain_of(lead["website"]) or ""
        preview = f"{site}/for/{make_slug(lead['business'], domain, settings['PROSPECTS_API_SECRET'])}?src=dashboard"
        score = f" and scored {lead['score']}/100 on mobile" if lead.get("score") else ""
        item = {"key": overrides.row_key({"Business": lead["business"], "Website": lead["website"]}), "sheet": inbound.SHEET,
                "business": lead["business"], "website": domain, "email": lead["email"], "contact": lead.get("name", ""),
                "phone": lead.get("phone", ""), "trade": lead.get("trade", ""), "area": lead.get("town", ""),
                "mobile_score": lead.get("score", ""), "preview": preview, "subject": "Your website report",
                "message": f"Please send me the full report on my website ({domain}) - I ran your free speed test{score}."}
        sid = self._remember("reply", item)
        self.start("all", {"sheet": inbound.SHEET}, then=(
            f"{lead['business']}'s preview: {preview.replace('src=dashboard', 'src=email')}\nThey asked for this - reply personally, today:",
            [[("✨ AI draft with their preview", f"ai:{sid}"), ("✏️ Write it", f"wr:{sid}")]]))

    def offer_draft(self, sid: str) -> None:
        item = self.items[sid]["item"]
        self.send(f"To {item.get('email')}:\n\n{self.drafts[sid]}",
                  [[("✅ Send it", f"sendreply:{sid}"), ("✏️ Edit", f"edit:{sid}"), ("🗑 Discard", f"discard:{sid}")]])

    def preview(self) -> None:
        import send_email

        out = send_email.preview_batch(self.panel.OUTREACH, env={**self.panel.load_settings()})
        if out.get("error"):
            self.send(out["error"])
            return
        emails = out.get("emails") or []
        for e in emails[:2]:
            self.send(f"To {e.get('business')} <{e.get('email')}> from {e.get('from', '')}\nSubject: {e.get('subject')}\n\n{e.get('body', '')}")
        skipped = [e for e in emails if e.get("skip")]
        self.send(f"{len(emails)} in the batch" + (f", {len(skipped)} would be skipped" if skipped else "") + " - the panel's Preview emails shows them all.")

    def start(self, action: str, extra: dict | None = None, then: tuple | None = None) -> None:
        p = self.panel
        body = {"action": action, "batch_size": BATCH_SIZE, "batch_scope": "all", "followup_size": BATCH_SIZE, "followup_days": 5,
                "send_gap": SEND_GAP_MIN, "send_from": "", **(extra or {})}
        if action not in ("batch", "followups", "send_batch", "send_followups", "autopilot_now", "all"):
            self.send("That can only be done on the panel.")
            return
        settings = p.load_settings()
        steps, error = p.build_steps(body, settings)
        if steps is None:
            self.send(error)
            return
        import autopilot

        if autopilot.is_running() and action != "autopilot_now":
            self.send("The autopilot is running, so this has to wait.", [[("⏹ Stop the autopilot", "apstop?")]])
            error = "busy"
        else:
            label = p.ACTIONS[action]
            error = p.JOB.start(label, steps, settings, keep_going=action in ("all", "batch"))
            if error:
                self.send(error)
        if error:
            if then:  # what to do next still applies - the run can happen later (🗂 Lists)
                self.send("Run the list later from 🗂 Lists. Meanwhile:\n" + then[0], then[1])
            return
        p.JOB.from_phone = True
        self.send(f"Started: {label}. I'll text you when it's done (📜 Log to look in on it).")
        threading.Thread(target=self._report_when_done, args=(then,), daemon=True).start()

    def _report_when_done(self, then: tuple | None = None) -> None:
        job = self.panel.JOB
        thread = job.thread
        if thread is not None:
            thread.join()
        state = job.state()
        tail = [ln.strip() for ln in state["lines"] if ln.strip() and not ln.startswith("> ")][-12:]
        self.send(f"{state['label']}:\n" + "\n".join(tail))
        if then:
            self.send(then[0], then[1])


_BOT: Bot | None = None


def start(panel) -> Bot | None:
    """Starts the bot thread when Telegram is set up; None otherwise. Called once by the panel."""
    global _BOT
    settings = panel.load_settings()
    token, chat = (settings.get("TELEGRAM_BOT_TOKEN") or "").strip(), (settings.get("TELEGRAM_CHAT_ID") or "").strip()
    if not token or not chat or _BOT is not None:
        return _BOT
    _BOT = Bot(panel, token, chat)
    threading.Thread(target=_BOT.run_forever, daemon=True).start()
    return _BOT
